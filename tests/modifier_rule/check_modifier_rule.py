"""Pass/fail checks for the Fried Specialty / modifier_rule issues #1-#6.

    python3 check_modifier_rule.py <bundle_dir> [--real-raw R85325_final_raw_gsheet.xlsx] [--keep DIR]

Builds the R99901 (DELTA) / R99902 (ABSOLUTE) fixtures, runs Part 1 and Part 2
end-to-end with the bundle's scripts, and checks the report AND the final raw.
The headline invariant: for every price the paper printed, what the customer is
charged in the final raw (dish_item + modifier_rule) equals the paper total.
--real-raw adds checks against a real wsheet export with a $0 dish whose price
lives on a required modifier group (R85325 "Seafood Boil").
Exit 1 if any check fails. Uses python3 + openpyxl directly (no .venv)."""
import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'lib'))
from xlformula import evaluate  # noqa: E402

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
PY = sys.executable
RESULTS = []


def check(cid, desc, ok, detail=''):
    RESULTS.append((cid, desc, bool(ok)))
    print(f"  {'PASS' if ok else 'FAIL'}  [{cid}] {desc}" + ('' if ok or not detail else f"  -- {detail}"))


def run(scripts, script, *args, log=None):
    p = subprocess.run([PY, os.path.join(scripts, script), *args], capture_output=True, text=True)
    if log:
        with open(log, 'w') as f:
            f.write(p.stdout + p.stderr)
    return p.returncode, p.stdout + p.stderr


def num(v):
    if v is None:
        return None
    try:
        return round(float(str(v).replace('$', '').replace(',', '').strip()), 2)
    except ValueError:
        return None


def fill_of(c):
    if c.fill is None or not c.fill.fill_type:
        return ''
    return {'ADD8E6': 'blue', 'FFB6C1': 'pink', 'C6EFCE': 'green',
            'E4DFEC': 'purple'}.get(c.fill.fgColor.rgb[-6:], c.fill.fgColor.rgb[-6:])


# ── readers ──────────────────────────────────────────────────────────────

def dish_item_cells(report):
    ws = load_workbook(report)['Dish Item']
    hdr = [str(c.value).strip() if c.value else '' for c in ws[1]]
    out = {}
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 2).value
        if not name:
            continue
        for ci, h in enumerate(hdr, 1):
            if h and h not in ('Dish Group', 'Dish Name:Primary', 'Dish Name:Secondary') \
                    and not h.startswith('_ref_'):
                c = ws.cell(r, ci)
                if c.value not in (None, ''):
                    out[(name, h)] = (num(c.value), fill_of(c))
    return out


def mrv_rows(report):
    """Modifier Rule tab → list of dicts keyed by header label."""
    wb = load_workbook(report)
    if 'Modifier Rule' not in wb.sheetnames:
        return []
    ws = wb['Modifier Rule']
    hdr_row, cols = None, {}
    for r in range(1, 12):
        labels = {str(ws.cell(r, c).value).strip(): c for c in range(1, 20) if ws.cell(r, c).value}
        if 'out_unit_price' in labels and 'modifier_item' in labels:
            hdr_row, cols = r, labels
            break
    rows = []
    ecol = next(c for k, c in cols.items() if k.startswith('Expected'))
    mcol = cols.get('Match?')
    for r in range(hdr_row + 1, ws.max_row + 1):
        mod = ws.cell(r, cols['modifier_item']).value
        if not mod:
            continue
        d = {k: ws.cell(r, c).value for k, c in cols.items()}
        d['_row'] = r
        # what the sheet SHOWS: 'by itself', Expected and Match? evaluated
        d['_by_itself'] = num(evaluate(wb, ws, d.get('by itself')))
        d['_expected'] = num(evaluate(wb, ws, ws.cell(r, ecol).value))
        d['_match'] = evaluate(wb, ws, ws.cell(r, mcol).value) if mcol else None
        d['_wb'], d['_ws'], d['_ecol'], d['_mcol'] = wb, ws, ecol, mcol
        rows.append(d)
    return rows


def match_if_price(d, price):
    """Match? the sheet would show if the Price cell held `price`."""
    ws, r = d['_ws'], d['_row']
    keep = ws.cell(r, 2).value
    ws.cell(r, 2).value = price
    try:
        return evaluate(d['_wb'], ws, ws.cell(r, d['_mcol']).value)
    finally:
        ws.cell(r, 2).value = keep


def mrv_get(rows, dish, mod, variant=None, group=None):
    for d in rows:
        if (d.get('dish_item') or '') == dish and d['modifier_item'] == mod \
                and (variant is None or d.get('variant') == variant) \
                and (group is None or (d.get('dish_group') or '') == group):
            return d
    return None


def sheet_text(report, tab):
    wb = load_workbook(report)
    if tab not in wb.sheetnames:
        return []
    return [[c for c in row if c not in (None, '')] for row in wb[tab].iter_rows(values_only=True)]


def propagation_sources(report):
    """(dish, modifier) → Source, from Fried Specialty's PROPAGATION DETAIL."""
    out, dish, on = {}, None, False
    for row in sheet_text(report, 'Fried Specialty'):
        if row and str(row[0]).startswith('PROPAGATION DETAIL'):
            on = True
            continue
        if not on or not row:
            continue
        if str(row[0]).startswith('⚠'):
            break
        if len(row) == 1:
            dish = str(row[0])
        elif len(row) >= 4 and row[0] != 'Old Price':
            out[(dish, str(row[1]))] = str(row[3])
    return out


def final_state(raw):
    """dish prices + modifier_rule rows from a raw gsheet (by header text)."""
    wb = load_workbook(raw)
    di = wb['dish_item']
    h3 = [str(c.value).strip() if c.value else '' for c in di[3]]
    h2 = [str(c.value).strip() if c.value else '' for c in di[2]]
    nc = h3.index('dish_item')
    pcols = {h3[i]: i for i in range(len(h3)) if 'unit_price' in h2[i]}
    dish = {}
    for row in di.iter_rows(min_row=4, values_only=True):
        if row[nc]:
            dish[row[nc]] = {v: num(row[i]) for v, i in pcols.items()}
    mr = wb['modifier_rule']
    m3 = [str(c.value).strip() if c.value else '' for c in mr[3]]
    ix = {k: m3.index(k) for k in ('out_unit_price', 'dish_group', 'dish_item', 'variant', 'modifier_item')}
    rules = {}
    for row in mr.iter_rows(min_row=4, values_only=True):
        key = (row[ix['dish_group']] or '', row[ix['dish_item']] or '', row[ix['variant']] or '',
               row[ix['modifier_item']] or '')
        rules[key] = num(row[ix['out_unit_price']])
    return dish, rules


def charged(dish, rules, d, variant, mod):
    base = dish.get(d, {}).get(variant) or 0.0
    rule = rules.get(('', d, variant, mod))
    return None if rule is None else round(base + rule, 2)


def tab_order_ok(report):
    names = load_workbook(report).sheetnames
    want = [t for t in ('Modifier Summary', 'Modifier Rule', 'Fried Specialty') if t in names]
    return [t for t in names if t in want] == want and len(want) == 3, names


# ── pipeline ─────────────────────────────────────────────────────────────

def part1(scripts):
    run(scripts, 'action1_raw_to_2a.py', 'raw.xlsx', '-o', 'preview.xlsx', log='log_a1.txt')
    run(scripts, 'ocr_sanity_check.py', 'ocr.json', log='log_sanity.txt')
    run(scripts, 'action2_ocr_to_2b_0625v2.py', 'ocr.json', 'preview.xlsx', '-o', 'ocr_out.xlsx',
        log='log_a2.txt')
    run(scripts, 'action3_compare_and_merge.py', 'preview.xlsx', 'ocr_out.xlsx', '-o', 'report.xlsx',
        log='log_a3.txt')
    run(scripts, 'price_inference.py', 'preview.xlsx', 'ocr_out.xlsx', '--report', 'report.xlsx',
        '--mode', 'apply', log='log_pi.txt')
    pre = run(scripts, 'modifier_rule_pricing.py', 'preflight', 'raw.xlsx', log='log_pre1.txt')
    run(scripts, 'modifier_rule_pricing.py', 'compare', 'ocr.json', 'raw.xlsx', '--report', 'report.xlsx',
        log='log_cmp.txt')
    return pre


def part2(scripts, report='report.xlsx', tag=''):
    run(scripts, 'action4_apply_modifier_changes.py', report, '-o', f'final_review{tag}.xlsx', log=f'log_a4{tag}.txt')
    a5 = run(scripts, 'action5_to_raw_gsheet.py', f'final_review{tag}.xlsx', 'raw.xlsx', '-o', f'final_raw{tag}.xlsx',
             log=f'log_a5{tag}.txt')
    # §2c runs on action5's output, so G11 sees the post-action5 dish_item state.
    pre = run(scripts, 'modifier_rule_pricing.py', 'preflight', f'final_raw{tag}.xlsx', log=f'log_pre2{tag}.txt')
    ap = run(scripts, 'modifier_rule_pricing.py', 'apply', report, f'final_raw{tag}.xlsx', '-o', f'final_raw{tag}.xlsx',
             log=f'log_apply{tag}.txt')
    return a5, pre, ap


def setup(work, rid):
    d = os.path.join(work, rid)
    os.makedirs(d, exist_ok=True)
    shutil.copy(os.path.join(work, f'{rid}_raw.xlsx'), os.path.join(d, 'raw.xlsx'))
    shutil.copy(os.path.join(work, f'{rid}_menu_ocr.json'), os.path.join(d, 'ocr.json'))
    os.chdir(d)
    return d


def edit_cell(path, tab, match, col_label, value):
    """Set <col_label> on the row of <tab> whose cells contain all of <match>."""
    wb = load_workbook(path)
    ws = wb[tab]
    hit = None
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if all(m in vals for m in match):
            hit = r
            break
    col = None
    for r in range(1, min(ws.max_row, 12) + 1):
        for c in range(1, ws.max_column + 1):
            if ws.cell(r, c).value == col_label:
                col = c
        if col:
            break
    ws.cell(hit, col).value = value
    wb.save(path)


# ── checks ───────────────────────────────────────────────────────────────

def check_delta(scripts, work):
    print('\nR99901 Delta Wok (DELTA model)')
    setup(work, 'R99901')
    pre_rc, pre_out = part1(scripts)
    rows = mrv_rows('report.xlsx')
    src = propagation_sources('report.xlsx')
    fs = sheet_text('report.xlsx', 'Fried Specialty')
    fs_flat = ' | '.join(str(c) for row in fs for c in row)

    # #2 — printed members keep their own value; unknowns not propagated from a split group
    pfr = mrv_get(rows, 'Chicken Wings', 'Pork Fried Rice')
    cfr = mrv_get(rows, 'Chicken Wings', 'Chicken Fried Rice')
    bfr = mrv_get(rows, 'Chicken Wings', 'Beef Fried Rice')
    check('#2', 'Wings + Pork Fried Rice keeps its printed +$2.20',
          pfr and num(pfr['out_unit_price']) == 2.20, pfr and pfr['out_unit_price'])
    check('#2', 'Wings + Chicken Fried Rice keeps its printed +$1.90',
          cfr and num(cfr['out_unit_price']) == 1.90, cfr and cfr['out_unit_price'])
    check('#2', 'unprinted Beef Fried Rice (same old-price group, printed siblings disagree) is NOT propagated',
          src.get(('Chicken Wings', 'Beef Fried Rice')) == 'NO DATA',
          src.get(('Chicken Wings', 'Beef Fried Rice')))
    col1 = [str(row[0]) if row and row[0] is not None else '' for row in fs]
    start = next((i for i, v in enumerate(col1) if v.startswith('⚠ CONFLICTS')), None)
    end = next((i for i, v in enumerate(col1) if start is not None and i > start and v == 'NOTES'), len(col1))
    gc = [v for v in col1[start + 1:end] if v.startswith('GROUP CONFLICT')] if start is not None else []
    check('#2', 'the split group is a callback item: GROUP CONFLICT on Chicken Wings, naming the unprinted '
                'Beef Fried Rice, under "⚠ CONFLICTS — NOT AUTO-RESOLVED" (not a note)',
          any('Chicken Wings' in g and 'Beef Fried Rice' in g.split('Unprinted', 1)[-1] for g in gc)
          and 'auto-picked' not in fs_flat, gc)
    check('#2', 'validation "Paper menu Total" shows the ACTUAL printed total ($10.19)',
          pfr and num(pfr.get('Paper menu Total')) == 10.19, pfr and pfr.get('Paper menu Total'))
    check('(e)', 'DELTA Expected = the live "by itself" (the Dish Item price) + the upcharge = the printed '
                 'total (Wings + Pork Fried Rice $10.19)',
          pfr and pfr['_expected'] == 10.19, pfr and pfr['_expected'])
    check('#2', 'derived (twin) row has no fake paper total',
          (lambda d: d is not None and d.get('Paper menu Total') in (None, ''))(
              mrv_get(rows, 'Fried Scallops', 'French Fries')))

    # #3 — "Plain Fried Rice" is an upcharge, not a By Itself anchor
    check('#3', 'preflight on a DELTA raw with "Plain Fried Rice" exits 0 / DELTA',
          pre_rc == 0 and 'DELTA' in pre_out, pre_out.strip().splitlines()[:2])
    pl = mrv_get(rows, 'Chicken Wings', 'Plain Fried Rice')
    check('#3', 'Plain Fried Rice updated as an upcharge (+$1.70)',
          pl and num(pl['out_unit_price']) == 1.70, pl and pl['out_unit_price'])

    # #5 — every multi-column section is processed
    bw_ff = mrv_get(rows, 'Buffalo Wings', 'French Fries')
    bw_pfr = mrv_get(rows, 'Buffalo Wings', 'Pork Fried Rice')
    check('#5', 'Buffalo Wings (2nd multi-column section) on the Modifier Rule tab at +$1.60 / +$2.20',
          bw_ff and bw_pfr and num(bw_ff['out_unit_price']) == 1.60 and num(bw_pfr['out_unit_price']) == 2.20,
          (bw_ff and bw_ff['out_unit_price'], bw_pfr and bw_pfr['out_unit_price']))
    cf_ff = mrv_get(rows, 'Chicken Fingers', 'French Fries')
    check('#5', 'Chicken Fingers in both sections with disagreeing fries totals: kept at $1.40 + flagged',
          cf_ff and num(cf_ff['out_unit_price']) == 1.40
          and src.get(('Chicken Fingers', 'French Fries')) == 'NO DATA'
          and any('Chicken Fingers' in str(c) and 'section' in str(c).lower() for row in fs for c in row),
          (cf_ff and cf_ff['out_unit_price'], src.get(('Chicken Fingers', 'French Fries'))))
    cf_pfr = mrv_get(rows, 'Chicken Fingers', 'Pork Fried Rice')
    check('#5', 'Chicken Fingers + Pork Fried Rice (both sections agree) still applied (+$2.20)',
          cf_pfr and num(cf_pfr['out_unit_price']) == 2.20, cf_pfr and cf_pfr['out_unit_price'])
    check('#5', 'Fried Specialty tab lists both sections',
          'Buffalo Wings' in fs_flat and 'Wings' in fs_flat and 'Specialties' in fs_flat)
    sides = sheet_text('ocr_out.xlsx', 'Fried Specialty Sides')
    check('#5', "action2 'Fried Specialty Sides' carries both sections",
          any('Buffalo Wings' in row for row in sides) and any('Chicken Wings' in row for row in sides))

    # #6 — tab order matches the doc
    ok, names = tab_order_ok('report.xlsx')
    check('#6', 'tab order ... Modifier Summary, Modifier Rule, Fried Specialty', ok, names)

    # #4 — specialist edits a group-level rule; Part 2 must write it
    edit_cell('report.xlsx', 'Modifier Rule', ['Fried Specialties', 'Extra Sauce'], 'out_unit_price', 0.75)
    a5, pre2, ap = part2(scripts)
    dish, rules = final_state('final_raw.xlsx')
    check('#4', 'group-level rule edit (Extra Sauce $0.50 → $0.75) reaches the final raw',
          rules.get(('Fried Specialties', '', 'Specialty', 'Extra Sauce')) == 0.75,
          rules.get(('Fried Specialties', '', 'Specialty', 'Extra Sauce')))
    check('#4', 'apply exits 0 with no drift on an unchanged raw', ap[0] == 0 and 'DRIFT' not in ap[1],
          ap[1][-300:])

    # headline invariant: charged == printed
    paper = [('Chicken Wings', 'Specialty', 'French Fries', 9.59),
             ('Chicken Wings', 'Specialty', 'Pork Fried Rice', 10.19),
             ('Chicken Wings', 'Specialty', 'Chicken Fried Rice', 9.89),
             ('Chicken Wings', 'Specialty', 'Plain Fried Rice', 9.69),
             ('Fried Shrimp', 'Specialty', 'French Fries', 11.55),
             ('Buffalo Wings', 'Order', 'French Fries', 10.55),
             ('Buffalo Wings', 'Order', 'Pork Fried Rice', 11.15)]
    bad = [(d, m, charged(dish, rules, d, v, m), t) for d, v, m, t in paper
           if charged(dish, rules, d, v, m) != t]
    check('#2/#5', 'final raw: every printed combo is charged exactly its paper total', not bad, bad)
    check('P2', 'Part 2 preflight on the post-action5 raw exits 0', pre2[0] == 0, pre2[1][:200])


def check_absolute(scripts, work):
    print('\nR99902 Absolute Wok (ABSOLUTE model)')
    setup(work, 'R99902')
    part1(scripts)
    di = dish_item_cells('report.xlsx')
    rows = mrv_rows('report.xlsx')
    q = sheet_text('report.xlsx', 'Questionnaire')

    cw = di.get(('Chicken Wings', 'Specialty'))
    check('#1', 'report: Chicken Wings dish_item stays $0.00 (paper Plain NOT written, not blue)',
          cw and cw[0] == 0.0 and cw[1] != 'blue', cw)
    sb = di.get(('Seafood Boil', 'Order'))
    check('#1', 'report: Seafood Boil ($0, price on a required modifier group) stays $0.00',
          sb and sb[0] == 0.0 and sb[1] != 'blue', sb)
    cf = di.get(('Chicken Fingers', 'Specialty'))
    check('#1', 'report: unprinted ABSOLUTE Chicken Fingers gets no inferred price (stays $0.00, not purple)',
          cf and cf[0] == 0.0 and cf[1] != 'purple', cf)
    p_abs = [row for row in q if row and str(row[0]).startswith('P-ABS-')]
    names = ' '.join(str(c) for row in p_abs for c in row)
    nums = {num(c) for row in p_abs for c in row if isinstance(c, (int, float)) or str(c).startswith('$')}
    check('#1', 'Questionnaire P-ABS lists Chicken Wings ($7.99) and Seafood Boil ($21.00)',
          'Chicken Wings' in names and 'Seafood Boil' in names and {7.99, 21.0} <= nums, p_abs)
    inf = [row for row in sheet_text('report.xlsx', 'Price Inference')
           if row and str(row[0]).startswith('INF-')
           and any(c in ('Chicken Fingers', 'Chicken Wings') for c in row)
           and any(c in ('APPLIED', 'SUGGESTED') for c in row)]
    check('#1', 'no inference applied or suggested for an ABSOLUTE $0.00 cell', not inf, inf)
    erp = di.get(('Egg Roll Platter', 'Specialty'))
    check('#1', 'regression: ordinary dish still updates (Egg Roll Platter $6.50 blue)',
          erp == (6.50, 'blue'), erp)
    fer = di.get(('Free Egg Roll', 'Order'))
    check('$0', 'freebie: $0.00 Free Egg Roll gets no inferred price although its whole group '
                'moved +$0.20 (T2)', fer and fer[0] == 0.0 and fer[1] != 'purple', fer)

    bi = mrv_get(rows, 'Chicken Wings', 'By Itself')
    ff = mrv_get(rows, 'Chicken Wings', 'French Fries')
    check('#1', 'Modifier Rule: By Itself rule carries the paper Plain ($7.99)',
          bi and num(bi['out_unit_price']) == 7.99, bi and bi['out_unit_price'])
    check('#1', 'Modifier Rule: fries rule carries the full paper total ($9.39)',
          ff and num(ff['out_unit_price']) == 9.39, ff and ff['out_unit_price'])
    check('#1', 'validation: Paper menu Total $9.39 (not $17.38) and Expected evaluates to $9.39',
          ff and num(ff.get('Paper menu Total')) == 9.39 and ff['_expected'] == 9.39,
          ff and (ff.get('Paper menu Total'), ff['_expected']))
    check('#1', 'validation: Match? TRUE as built; a wrong edit (8.90) shows Match? FALSE — evaluated, '
                'not read off the formula text',
          ff and ff['_match'] == 'TRUE' and match_if_price(ff, 8.90) == 'FALSE',
          ff and (ff['_match'], match_if_price(ff, 8.90)))
    ok, names_ = tab_order_ok('report.xlsx')
    check('#6', 'tab order ... Modifier Summary, Modifier Rule, Fried Specialty', ok, names_)

    a5, pre2, ap = part2(scripts)
    dish, rules = final_state('final_raw.xlsx')
    check('#1', 'final raw: Chicken Wings dish_item $0.00; Seafood Boil dish_item $0.00',
          dish['Chicken Wings']['Specialty'] == 0.0 and dish['Seafood Boil']['Order'] == 0.0,
          (dish['Chicken Wings'], dish['Seafood Boil']))
    got = {m: charged(dish, rules, 'Chicken Wings', 'Specialty', m)
           for m in ('By Itself', 'French Fries', 'Pork Fried Rice')}
    check('#1', 'final raw: customer is charged the paper totals (Plain $7.99, +fries $9.39, +PFR $9.89)',
          got == {'By Itself': 7.99, 'French Fries': 9.39, 'Pork Fried Rice': 9.89}, got)
    check('#1', 'Part 2 preflight on the post-action5 raw exits 0 / ABSOLUTE',
          pre2[0] == 0 and 'ABSOLUTE' in pre2[1], pre2[1][:200])
    check('#1', 'apply on the post-action5 raw exits 0', ap[0] == 0, ap[1][-300:])

    # Guard 1: a specialist hand-types $7.99 on the $0 dish → action5 must not write it.
    shutil.copy('report.xlsx', 'report_edit.xlsx')
    edit_cell('report_edit.xlsx', 'Dish Item', ['Chicken Wings'], 'Specialty', 7.99)
    a5e, pre2e, ape = part2(scripts, 'report_edit.xlsx', tag='_edit')
    dish_e, _ = final_state('final_raw_edit.xlsx')
    check('#1', 'action5 refuses to write a price onto a $0.00 dish whose price lives on modifier rules',
          dish_e['Chicken Wings']['Specialty'] == 0.0 and 'ABSOLUTE' in a5e[1],
          (dish_e['Chicken Wings'], a5e[1][-400:]))

    # Guard 2: apply must refuse a DOUBLE_CHARGE state (dish_item priced + By Itself priced).
    wb = load_workbook('raw.xlsx')
    ws = wb['dish_item']
    hdr = [c.value for c in ws[3]]
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, hdr.index('dish_item') + 1).value == 'Chicken Wings':
            ws.cell(r, hdr.index('Specialty') + 1).value = '$7.99'
    wb.save('raw_double.xlsx')
    rc, out = run(scripts, 'modifier_rule_pricing.py', 'apply', 'report.xlsx', 'raw_double.xlsx',
                  '-o', 'final_double.xlsx', log='log_apply_double.txt')
    check('#1', 'apply refuses a DOUBLE_CHARGE raw (exit 1, nothing written)',
          rc == 1 and not os.path.exists('final_double.xlsx'), (rc, out[-300:]))


def check_village(scripts, work):
    print('\nR99904 Village Wok (China Village shape: model differs per dish AND per size)')
    setup(work, 'R99904')
    pre_rc, pre_out = part1(scripts)
    check('CV', 'preflight classifies the restaurant (MIXED, exit 0) — the restaurant-wide '
                '"Plain $0.00" rule is not taken as a dish anchor',
          pre_rc == 0 and 'MIXED' in pre_out, pre_out.strip().splitlines()[:2])
    wb = load_workbook('preview.xlsx')
    held = sorted((r[1], r[2]) for r in wb['absolute_pricing'].iter_rows(min_row=2, values_only=True)
                  if r[1]) if 'absolute_pricing' in wb.sheetnames else []
    check('CV', 'absolute_pricing = the $0 cells only (not Cut Wings 50 Pieces, not Bento Box)',
          held == [('Barbecue Boneless Spare Ribs', 'Specialty'), ('Chicken Nuggets', 'Specialty'),
                   ('Cut Wings', '10 Pieces'), ('Cut Wings', '5 Pieces'),
                   ('Fried Fish', 'Specialty'), ('Fried Scallop', 'Specialty'),
                   ('Whole Wing', '3 Pieces'), ('Whole Wing', '4 Pieces')], held)
    di = dish_item_cells('report.xlsx')
    check('CV', 'Dish Item: ABSOLUTE cells stay $0.00; DELTA sizes take the paper price (Cut Wings 50 pcs $52.00)',
          di.get(('Chicken Nuggets', 'Specialty'), (None,))[0] == 0.0
          and di.get(('Cut Wings', '5 Pieces'), (None,))[0] == 0.0
          and di.get(('Cut Wings', '50 Pieces')) == (52.0, 'blue'),
          (di.get(('Chicken Nuggets', 'Specialty')), di.get(('Cut Wings', '5 Pieces')),
           di.get(('Cut Wings', '50 Pieces'))))
    rows = mrv_rows('report.xlsx')

    def price(dish, var, mod):
        d = mrv_get(rows, dish, mod, variant=var)
        return num(d['out_unit_price']) if d else None
    got = {k: price(*k) for k in [
        ('Chicken Nuggets', 'Specialty', 'By Itself'), ('Chicken Nuggets', 'Specialty', 'French Fries'),
        ('Cut Wings', '5 Pieces', 'By Itself'), ('Cut Wings', '5 Pieces', 'French Fries'),
        ('Cut Wings', '50 Pieces', 'French Fries'), ('Cut Wings', '50 Pieces', 'By Itself'),
        ('Whole Wing', '4 Pieces', 'By Itself'), ('Whole Wing', '4 Pieces', 'French Fries')]}
    want = {('Chicken Nuggets', 'Specialty', 'By Itself'): 7.25,     # ABSOLUTE: Plain → By Itself
            ('Chicken Nuggets', 'Specialty', 'French Fries'): 11.25, # ABSOLUTE: full total
            ('Cut Wings', '5 Pieces', 'By Itself'): 6.75,
            ('Cut Wings', '5 Pieces', 'French Fries'): 10.45,
            ('Cut Wings', '50 Pieces', 'French Fries'): 3.50,        # DELTA size: 55.50 − 52.00
            ('Cut Wings', '50 Pieces', 'By Itself'): 0.00,           # DELTA: Plain went to dish_item
            ('Whole Wing', '4 Pieces', 'By Itself'): 9.75,           # size read from "4pcs"
            ('Whole Wing', '4 Pieces', 'French Fries'): 13.45}
    bad = {k: (got[k], v) for k, v in want.items() if got[k] != v}
    check('CV', 'Modifier Rule: each size priced by its own model (ABSOLUTE totals / DELTA upcharges)',
          not bad, bad)
    unprinted = [d for d in rows if (d.get('dish_item'), d.get('variant')) in
                 (('Cut Wings', '10 Pieces'), ('Whole Wing', '3 Pieces'))]
    raw_old = {}
    for r in load_workbook('raw.xlsx')['modifier_rule'].iter_rows(min_row=4, values_only=True):
        raw_old[(r[9], r[10], r[12])] = num(r[2])
    moved = [(d['dish_item'], d['variant'], d['modifier_item'], num(d['out_unit_price']))
             for d in unprinted
             if num(d['out_unit_price']) != raw_old.get((d['dish_item'], d['variant'], d['modifier_item']))]
    check('CV', 'sizes the paper did not print keep their OWN prices (no collapse across sizes)',
          unprinted and not moved, moved)
    fs = ' | '.join(str(c) for row in sheet_text('report.xlsx', 'Fried Specialty') for c in row)
    check('CV', 'a paper row with no size for a sized dish is flagged, not applied (Whole Wing)',
          'VARIANT UNKNOWN' in fs and 'Whole Wing' in fs)
    # China Village layout: the combo box prints only the with-side price
    got = {k: price(*k) for k in [('Fried Scallop', 'Specialty', 'By Itself'),
                                  ('Fried Scallop', 'Specialty', 'French Fries'),
                                  ('Chicken Nuggets', 'Specialty', 'Plain Lo Mein'),
                                  ('Spare Rib Tips', 'Order', 'French Fries')]}
    want = {('Fried Scallop', 'Specialty', 'By Itself'): 7.25,        # Plain joined from APPETIZER
            ('Fried Scallop', 'Specialty', 'French Fries'): 11.25,    # combo box total
            ('Chicken Nuggets', 'Specialty', 'Plain Lo Mein'): 12.25, # bare "Lo Mein" column
            ('Spare Rib Tips', 'Order', 'French Fries'): 3.40}        # DELTA, no plain: not applied
    bad = {k: (got[k], v) for k, v in want.items() if got[k] != v}
    check('CV', 'combo box: ABSOLUTE rows take the full total + the Plain from APPETIZER; a DELTA '
                'row without a Plain is not applied; "Lo Mein" column maps', not bad, bad)
    check('CV', 'combo box: flagged NO PLAIN for the DELTA dish; the join is noted',
          'NO PLAIN' in fs and 'Spare Rib Tips' in fs and 'PLAIN FROM ANOTHER SECTION' in fs)
    bl = price('Barbecue Boneless Spare Ribs', 'Specialty', 'French Fries')
    check('CV', 'a combo row for a dish without rules ("BBQ Spare Ribs") does not land on a '
                'similar-named dish (Boneless takes its own $11.45, no conflict)',
          bl == 11.45 and 'Barbecue Boneless Spare Ribs / French Fries' not in fs, bl)
    q = sheet_text('report.xlsx', 'Questionnaire')
    pabs = [row for row in q if row and str(row[0]).startswith('P-ABS-')
            and any(str(c).startswith('Fried Scallop') for c in row)]
    nums = {num(c) for row in pabs for c in row if isinstance(c, (int, float))}
    di2 = dish_item_cells('report.xlsx')
    check('CV', 'the combo price never lands on a dish price (P-ABS shows the Plain $7.25; '
                'Spare Rib Tips Order stays $7.55)',
          pabs and 7.25 in nums and 11.25 not in nums
          and di2.get(('Spare Rib Tips', 'Order'), (None,))[0] == 7.55,
          (pabs, di2.get(('Spare Rib Tips', 'Order'))))

    a5, pre2, ap = part2(scripts)
    dish, rules = final_state('final_raw.xlsx')
    paper = [('Chicken Nuggets', 'Specialty', 'French Fries', 11.25), ('Chicken Nuggets', 'Specialty', 'By Itself', 7.25),
             ('Fried Fish', 'Specialty', 'Fried Rice', 11.25),
             ('Cut Wings', '5 Pieces', 'French Fries', 10.45), ('Cut Wings', '5 Pieces', 'By Itself', 6.75),
             ('Cut Wings', '50 Pieces', 'French Fries', 55.50), ('Cut Wings', '50 Pieces', 'Fried Rice', 55.50),
             ('Whole Wing', '4 Pieces', 'French Fries', 13.45),
             ('Fried Scallop', 'Specialty', 'By Itself', 7.25), ('Fried Scallop', 'Specialty', 'Fried Rice', 11.25)]
    bad = [(d, v, m, charged(dish, rules, d, v, m), t) for d, v, m, t in paper if charged(dish, rules, d, v, m) != t]
    check('CV', 'final raw: every printed combo, every size, is charged exactly its paper total', not bad, bad)
    check('CV', 'final raw: unprinted sizes charged as before (Cut Wings 10 pcs + fries $13.95)',
          charged(dish, rules, 'Cut Wings', '10 Pieces', 'French Fries') == 13.95)
    check('CV', 'Part 2 preflight + apply on the post-action5 raw exit 0',
          pre2[0] == 0 and ap[0] == 0, (pre2[1][:160], ap[1][-200:]))
    wb = load_workbook('raw.xlsx')
    ws = wb['dish_item']
    hdr = [c.value for c in ws[3]]
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, hdr.index('dish_item') + 1).value == 'Chicken Nuggets':
            ws.cell(r, hdr.index('Specialty') + 1).value = '$7.25'
    wb.save('raw_double.xlsx')
    rc, out = run(scripts, 'modifier_rule_pricing.py', 'preflight', 'raw_double.xlsx', log='log_pre_double.txt')
    check('CV', 'a dish priced on BOTH dish_item and its By Itself rule still blocks (DOUBLE_CHARGE)',
          rc == 1 and 'DOUBLE' in out, out[:200])


def check_real(scripts, work, real_raw):
    print(f'\nReal raw: {os.path.basename(real_raw)}')
    d = os.path.join(work, 'real')
    os.makedirs(d, exist_ok=True)
    os.chdir(d)
    shutil.copy(real_raw, 'raw.xlsx')
    rc, out = run(scripts, 'action1_raw_to_2a.py', 'raw.xlsx', '-o', 'preview.xlsx', log='log_a1.txt')
    wb = load_workbook('preview.xlsx')
    held = []
    if 'absolute_pricing' in wb.sheetnames:
        held = [(r[1], r[2]) for r in wb['absolute_pricing'].iter_rows(min_row=2, values_only=True) if r[1]]
    check('R', 'action1 flags exactly Seafood Boil/Order as "$0, price on modifiers"',
          held == [('Seafood Boil', 'Order')], held)
    rc, out = run(scripts, 'modifier_rule_pricing.py', 'preflight', 'raw.xlsx', log='log_pre.txt')
    check('R', 'preflight on the real raw exits 0 (DELTA)', rc == 0 and 'DELTA' in out, out[:200])
    ocr = {'restaurant': {'name': 'real'}, 'sections': [
        {'menu_section': 'Seafood Boil', 'price_columns': ['Order'], 'items': [
            {'name': 'Seafood Boil', 'prices': {'Order': 21.00}}]},
        {'menu_section': 'Fried Specialties', 'price_columns': ['Order'], 'items': [
            {'name': 'Buffalo Wing', 'prices': {'Order': 14.49}}]}]}
    with open('ocr.json', 'w') as f:
        json.dump(ocr, f)
    run(scripts, 'action2_ocr_to_2b_0625v2.py', 'ocr.json', 'preview.xlsx', '-o', 'ocr_out.xlsx', log='log_a2.txt')
    run(scripts, 'action3_compare_and_merge.py', 'preview.xlsx', 'ocr_out.xlsx', '-o', 'report.xlsx',
        log='log_a3.txt')
    di = dish_item_cells('report.xlsx')
    check('R', 'Seafood Boil paper $21 held (stays $0.00); Buffalo Wing updated to $14.49',
          di.get(('Seafood Boil', 'Order'), (None,))[0] == 0.0
          and di.get(('Buffalo Wing', 'Order'), (None,))[0] == 14.49,
          (di.get(('Seafood Boil', 'Order')), di.get(('Buffalo Wing', 'Order'))))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('bundle', help='bundle root (the dir holding menu-change-payload/)')
    ap.add_argument('--real-raw', help='optional real raw gsheet with a $0 modifier-priced dish (R85325)')
    ap.add_argument('--keep', help='work dir to keep (default: temp, deleted)')
    a = ap.parse_args()
    real_raw = os.path.abspath(a.real_raw) if a.real_raw else None
    scripts = os.path.join(os.path.abspath(a.bundle), 'menu-change-payload', 'scripts')
    work = os.path.abspath(a.keep) if a.keep else tempfile.mkdtemp(prefix='mr_checks_')
    if a.keep and os.path.isdir(work):
        shutil.rmtree(work)
    os.makedirs(work, exist_ok=True)
    os.chdir(work)
    subprocess.run([PY, os.path.join(HERE, 'build_fixture.py')], check=True, capture_output=True)
    check_delta(scripts, work)
    check_absolute(scripts, work)
    check_village(scripts, work)
    if real_raw:
        check_real(scripts, work, real_raw)
    os.chdir(HERE)
    if not a.keep:
        shutil.rmtree(work, ignore_errors=True)
    failed = [r for r in RESULTS if not r[2]]
    print(f"\nmodifier_rule: {len(RESULTS) - len(failed)}/{len(RESULTS)} PASS")
    sys.exit(1 if failed else 0)


if __name__ == '__main__':
    main()
