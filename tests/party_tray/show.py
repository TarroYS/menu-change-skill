"""Print Party Tray column at every stage: system, OCR out, report (with fill), final raw."""
import sys
from openpyxl import load_workbook
d = sys.argv[1]
def col(ws, hdr_row, name):
    for c in ws[hdr_row]:
        if str(c.value).strip() == name: return c.column
FILLS = {'ADD8E6':'blue','FFB6C1':'pink','C6EFCE':'green','E4DFEC':'PURPLE'}
pv = load_workbook(f'{d}/preview.xlsx')['price']; ob = load_workbook(f'{d}/ocr_out.xlsx')['price']
rp = load_workbook(f'{d}/report.xlsx')['Dish Item']; fr = load_workbook(f'{d}/final_raw.xlsx')['dish_item']
pc, oc, rc = col(pv,1,'Party Tray'), col(ob,1,'Party Tray'), col(rp,1,'Party Tray')
lc_rp = col(rp,1,'Large'); fc = col(fr,3,'Party Tray'); fn = col(fr,3,'dish_item')
final = {fr.cell(r,fn).value: fr.cell(r,fc).value for r in range(4, fr.max_row+1)}
print(f"{'Dish':24} {'Large(rpt)':>10} {'Tray sys':>9} {'Tray OCR':>9} {'Tray rpt':>9} {'fill':7} {'FINAL RAW':>9}")
for r in range(2, pv.max_row+1):
    name = pv.cell(r,2).value
    if not name: continue
    c = rp.cell(r,rc); f = c.fill.fgColor.rgb[-6:] if c.fill and c.fill.fill_type else ''
    print(f"{name:24} {str(rp.cell(r,lc_rp).value or ''):>10} {str(pv.cell(r,pc).value or ''):>9} {str(ob.cell(r,oc).value or ''):>9} {str(c.value or ''):>9} {FILLS.get(f,''):7} {str(final.get(name) or ''):>9}")
