"""Pass/fail checks for Party Tray (v4.3), on the harness fixture.

    python3 check_party_tray.py <bundle_dir>

Runs build_fixture.py + run_p1p2.sh (normal and 'mixed' fixtures) against the
bundle and asserts the v4.3 behaviour show.py prints: printed trays compared,
unprinted trays derived purple from Large × 5 − $2.00 only where the row
follows the formula, outliers / decreases / unprinted bases left for the
callback, and no formula at all on the mixed fixture. Exit 1 on any failure."""
import os
import re
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
RESULTS = []
FILLS = {'ADD8E6': 'blue', 'FFB6C1': 'pink', 'C6EFCE': 'green', 'E4DFEC': 'purple'}


def check(desc, ok, detail=''):
    RESULTS.append(bool(ok))
    print(f"  {'PASS' if ok else 'FAIL'}  {desc}" + ('' if ok or not detail else f'  -- {detail}'))


def num(v):
    try:
        return round(float(str(v).replace('$', '').strip()), 2)
    except ValueError:
        return None


def trays(out):
    """dish → (report tray value, fill, final raw tray value), prices as numbers."""
    rp = load_workbook(os.path.join(out, 'report.xlsx'))['Dish Item']
    hdr = [c.value for c in rp[1]]
    tc = hdr.index('Party Tray') + 1
    fr = load_workbook(os.path.join(out, 'final_raw.xlsx'))['dish_item']
    h3 = [c.value for c in fr[3]]
    final = {fr.cell(r, h3.index('dish_item') + 1).value: fr.cell(r, h3.index('Party Tray') + 1).value
             for r in range(4, fr.max_row + 1)}
    res = {}
    for r in range(2, rp.max_row + 1):
        name = rp.cell(r, 2).value
        if not name:
            continue
        c = rp.cell(r, tc)
        f = FILLS.get(c.fill.fgColor.rgb[-6:], '') if c.fill and c.fill.fill_type else ''
        res[name] = (num(c.value) if c.value not in (None, '') else None, f,
                     num(final.get(name)) if final.get(name) not in (None, '') else None)
    return res


def final_strings(out):
    """Party Tray cells of the final raw, verbatim (to check the $X.XX format)."""
    fr = load_workbook(os.path.join(out, 'final_raw.xlsx'))['dish_item']
    h3 = [c.value for c in fr[3]]
    return [fr.cell(r, h3.index('Party Tray') + 1).value for r in range(4, fr.max_row + 1)
            if fr.cell(r, h3.index('Party Tray') + 1).value not in (None, '')]


def questionnaire(out):
    ws = load_workbook(os.path.join(out, 'report.xlsx'))['Questionnaire']
    return [' '.join(str(v) for v in row if v not in (None, '')) for row in ws.iter_rows(values_only=True)]


def main():
    bundle = os.path.abspath(sys.argv[1])
    work = tempfile.mkdtemp(prefix='pt_checks_')
    for f in ('build_fixture.py', 'run_p1p2.sh'):
        shutil.copy(os.path.join(HERE, f), work)
    env = dict(os.environ, PATH=os.path.dirname(sys.executable) + os.pathsep + os.environ.get('PATH', ''))
    run = lambda *a: subprocess.run(list(a), cwd=work, env=env, capture_output=True, text=True)
    run(sys.executable, 'build_fixture.py')
    run(sys.executable, 'build_fixture.py', 'mixed')
    p = run('bash', 'run_p1p2.sh', bundle, 'R99999_raw.xlsx', 'out')
    check('pipeline runs (action1 → … → action5)', p.returncode == 0, p.stdout[-300:])
    if p.returncode:
        sys.exit(1)
    t = trays(os.path.join(work, 'out'))
    print('\nR99999 (Party Tray = Large × 5 − $2.00)')
    check('printed tray compared like any variant (Sesame Chicken $68.00 blue → final)',
          t['Sesame Chicken'] == (68.0, 'blue', 68.0), t['Sesame Chicken'])
    for dish, want in (("General Tso's Chicken", 67.75), ('Chicken with Broccoli', 62.75)):
        check(f'unprinted tray derived from the printed Large, purple ({dish} ${want:.2f})',
              t[dish] == (want, 'purple', want), t[dish])
    # formulas are per dish group (maintainer, 2026-10-05): Beef has 3 trays, one an outlier —
    # 2/3 is too few to trust a formula for Beef, so its changed trays are asked, not derived
    for dish, old in (('Pepper Steak', 72.75), ('Mongolian Beef', 75.50)):
        check(f'Beef group cannot confirm a formula (2 of 3 trays fit) → {dish} tray not derived '
              f'(stays ${old:.2f})', t[dish] == (old, '', old), t[dish])
    check('formula outlier not derived (Beef with Broccoli stays $80.00)',
          t['Beef with Broccoli'] == (80.0, '', 80.0), t['Beef with Broccoli'])
    check('unchanged base → tray unchanged (Kung Pao Chicken $57.75)',
          t['Kung Pao Chicken'] == (57.75, '', 57.75), t['Kung Pao Chicken'])
    check('decreased base → tray not written (Lemon Chicken $60.50)',
          t['Lemon Chicken'] == (60.5, '', 60.5), t['Lemon Chicken'])
    check('dish not on paper → tray unchanged (Hunan Chicken $62.75)',
          t['Hunan Chicken'] == (62.75, '', 62.75), t['Hunan Chicken'])
    check('tray-only dish printed → blue (Lo Mein Tray $48.00)',
          t['Lo Mein Tray'] == (48.0, 'blue', 48.0), t['Lo Mein Tray'])
    check('tray-only dish unprinted → unchanged (Fried Rice Tray $45.00)',
          t['Fried Rice Tray'] == (45.0, '', 45.0), t['Fried Rice Tray'])
    fs = final_strings(os.path.join(work, 'out'))
    check('final raw keeps the $X.XX format on every tray price (written ones included)',
          fs and all(isinstance(x, str) and x.startswith('$') and len(x.split('.')[-1]) == 2 for x in fs),
          [x for x in fs if not (isinstance(x, str) and x.startswith('$'))])
    q = questionnaire(os.path.join(work, 'out'))
    check('P-TRAY-0 states the detected formula',
          any(r.startswith('P-TRAY-0') and 'Large × 5' in r and '2.00' in r for r in q),
          [r for r in q if r.startswith('P-TRAY')][:2])
    check('outlier asked on the callback (P-TRAY row for Beef with Broccoli)',
          any(r.startswith('P-TRAY-') and 'Beef with Broccoli' in r for r in q))
    check('… and so are the Beef trays with no group formula (Pepper Steak, Mongolian Beef); P-TRAY-0 '
          'says why for Beef', all(any(r.startswith('P-TRAY-') and d in r for r in q)
                                  for d in ('Pepper Steak', 'Mongolian Beef'))
          and any(r.startswith('P-TRAY-0') and 'No reliable tray formula for Beef' in r for r in q),
          [r for r in q if r.startswith('P-TRAY')][:6])
    check('P1-TRAY row asks for Party Tray prices (the restaurant has trays); no old A-1 row',
          any(r.startswith('P1-TRAY') and 'Party Tray pricing' in r for r in q)
          and not any(r.startswith('A-1') for r in q))

    print('\nR99999 mixed (no dominant tray ratio)')
    p = run('bash', 'run_p1p2.sh', bundle, 'R99999_raw_mixed.xlsx', 'out_mixed')
    check('pipeline runs', p.returncode == 0, p.stdout[-300:])
    if p.returncode == 0:
        tm = trays(os.path.join(work, 'out_mixed'))
        check('no formula → no tray derived (no purple)', not any(v[1] == 'purple' for v in tm.values()),
              {k: v for k, v in tm.items() if v[1] == 'purple'})
        qm = questionnaire(os.path.join(work, 'out_mixed'))
        check('P-TRAY-0 says why no formula was used (no reliable formula + the agreement it found)',
              any(r.startswith('P-TRAY-0') and 'No reliable tray formula' in r and re.search(r'\d+/\d+', r)
                  for r in qm), [r for r in qm if r.startswith('P-TRAY')][:1])

    shutil.rmtree(work, ignore_errors=True)
    print(f'\nparty_tray: {sum(RESULTS)}/{len(RESULTS)} PASS')
    sys.exit(0 if all(RESULTS) else 1)


if __name__ == '__main__':
    main()
