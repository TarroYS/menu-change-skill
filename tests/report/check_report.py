"""Pass/fail checks for the specialist's edit sheet and how Part 2 reads it (2026-09-30).

    python3 check_report.py <bundle_dir> [--real-report <an edited pre-2026-09-30 report>]

Part 1 (the report the specialist edits):
  #1  every price is a plain number — no '$X.XX' strings in any tab
  #8  a P1-TRAY row asks for Party Tray prices when the restaurant has trays
  #9  no P3 / Additional Questions / P5 / P6
  #7  column I "Notes → Wonka" exists and Part 1 leaves it empty
  P-ABS rows name the exact tab / row / column to edit (and link to it)
  x / X over a Dish Item price shows in its own colour (ahead of the apricot rule)
Part 2 (a specialist-edited copy, through the wsheet TSV bridge like ACE):
  #1  '$9.25' typed the old way and 8.95 typed plain both apply; the raw keeps $X.XX;
      0 is a real price (Modifier Summary 0.50 → 0)
  #3  x over a price clears it in the raw AND in the TSV; x on an empty cell is ignored
  #2  strikethrough on a dish name (whole cell or part) → handed to Wonka
  #7  every Notes → Wonka note → handed to Wonka with its row's context;
      an older report's Notes column (script text) is not misread
Exit 1 if any check fails."""
import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
RESULTS = []
MONEY = re.compile(r'^\s*\$\s*[+-]?\d')
NOTES_HEADER = 'Notes → Wonka (write an instruction)'


def check(desc, ok, detail=''):
    RESULTS.append(bool(ok))
    print(f"  {'PASS' if ok else 'FAIL'}  {desc}" + ('' if ok or not detail else f'  -- {detail}'))


def run(scripts, script, *args):
    p = subprocess.run([sys.executable, os.path.join(scripts, script), *args],
                       capture_output=True, text=True)
    return p.returncode, p.stdout + p.stderr


def part1(scripts, has_sides):
    run(scripts, 'action1_raw_to_2a.py', 'raw.xlsx', '-o', 'preview.xlsx')
    run(scripts, 'action2_ocr_to_2b_0625v2.py', 'ocr.json', 'preview.xlsx', '-o', 'ocr_out.xlsx')
    run(scripts, 'action3_compare_and_merge.py', 'preview.xlsx', 'ocr_out.xlsx', '-o', 'report.xlsx')
    run(scripts, 'price_inference.py', 'preview.xlsx', 'ocr_out.xlsx', '--report', 'report.xlsx',
        '--mode', 'apply')
    if has_sides:
        run(scripts, 'modifier_rule_pricing.py', 'compare', 'ocr.json', 'raw.xlsx', '--report', 'report.xlsx')


def fixture(work, name, builder_dir, builder_args, raw, ocr):
    d = os.path.join(work, name)
    os.makedirs(d, exist_ok=True)
    os.chdir(d)
    subprocess.run([sys.executable, os.path.join(HERE, '..', builder_dir, 'build_fixture.py'), *builder_args],
                   check=True, capture_output=True)
    shutil.copy(raw, 'raw.xlsx')
    shutil.copy(ocr, 'ocr.json')
    return d


# ── Part 1 checks ────────────────────────────────────────────────────────

def check_numeric(tag):
    wb = load_workbook('report.xlsx')
    bad, qbad, fmt_bad = [], [], []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and MONEY.match(c.value):
                    bad.append(f'{ws.title}!{c.coordinate}={c.value}')
                if (ws.title == 'Questionnaire' and c.column in (5, 6, 7)
                        and isinstance(c.value, str) and '$' in c.value):
                    qbad.append(f'{c.coordinate}={c.value}')
    di = wb['Dish Item']
    hdr = [c.value for c in di[1]]
    names = {'Dish Group', 'Dish Name:Primary', 'Dish Name:Secondary'}
    for ci, h in enumerate(hdr, 1):
        if not h or h in names or str(h).startswith('_ref_'):
            continue
        for r in range(2, di.max_row + 1):
            c = di.cell(r, ci)
            if c.value not in (None, '') and (not isinstance(c.value, (int, float))
                                              or c.number_format != '0.00'):
                fmt_bad.append(f'{c.coordinate}={c.value!r}/{c.number_format}')
    check(f'{tag}: no "$X.XX" price strings in any tab of the report', not bad, bad[:6])
    check(f'{tag}: Questionnaire price columns carry no "$"', not qbad, qbad[:6])
    check(f'{tag}: Dish Item prices are numbers shown as 0.00', not fmt_bad, fmt_bad[:6])


def check_questionnaire(tag, expect_tray):
    ws = load_workbook('report.xlsx')['Questionnaire']
    hdr = [c.value for c in ws[1]]
    check(f'{tag}: Questionnaire columns end with Details + "{NOTES_HEADER}"',
          hdr[7:9] == ['Details', NOTES_HEADER], hdr)
    filled = [r for r in range(2, ws.max_row + 1) if ws.cell(r, 9).value not in (None, '')]
    check(f'{tag}: Part 1 writes nothing in the Notes → Wonka column', not filled, filled[:5])
    col_a = [str(ws.cell(r, 1).value) for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
    gone = [x for x in col_a if x.startswith(('P3', 'P5', 'P6', 'A-', 'ADDITIONAL'))]
    check(f'{tag}: no P3 / Additional Questions / P5 / P6 rows', not gone, gone[:5])
    has_tray = any(x == 'P1-TRAY' for x in col_a)
    check(f'{tag}: P1-TRAY row {"present" if expect_tray else "absent"} '
          f'({"has" if expect_tray else "no"} Party Tray prices)', has_tray == expect_tray, col_a[:4])
    if expect_tray:
        p1 = col_a.index('P1 — ON MENU, NOT MATCHED TO SYSTEM') if 'P1 — ON MENU, NOT MATCHED TO SYSTEM' in col_a else -1
        check(f'{tag}: P1-TRAY sits in the P1 section', p1 >= 0 and col_a[p1 + 1] == 'P1-TRAY', col_a[:4])


def check_pabs_pointers():
    wb = load_workbook('report.xlsx')
    q, mr, ms = wb['Questionnaire'], wb['Modifier Rule'], wb['Modifier Summary']
    rows = {str(q.cell(r, 4).value): q.cell(r, 8) for r in range(2, q.max_row + 1)
            if str(q.cell(r, 1).value or '').startswith('P-ABS-')}
    rule = rows.get('Chicken Wings [Specialty]')
    m = re.search(r"'Modifier Rule' tab · row (\d+) · column B", str(rule.value if rule else ''))
    target = [mr.cell(int(m.group(1)), c).value for c in (4, 5, 6)] if m else None
    check("P-ABS (modifier rules) names the exact Modifier Rule row + column B, and that row is "
          "Chicken Wings' By Itself", target == ['Chicken Wings', 'Specialty', 'By Itself'], (rule.value if rule else None, target))
    check('  … and links to that cell',
          m and rule.hyperlink is not None and rule.hyperlink.location == f"'Modifier Rule'!B{m.group(1)}",
          rule.hyperlink.location if rule is not None and rule.hyperlink else None)
    choice = rows.get('Seafood Boil [Order]')
    m2 = re.search(r"'Modifier Summary' tab · rows? (\d+)(?:–(\d+))? \(group 'Select Seafood'\) · "
                   r"column 'Order' \((\w+)\)", str(choice.value if choice else ''))
    ok = False
    if m2:
        a, b, col = int(m2.group(1)), int(m2.group(2) or m2.group(1)), m2.group(3)
        hdr = {c.column_letter: c.value for c in ms[1]}
        ok = hdr.get(col) == 'Order' and all(ms.cell(r, 1).value == 'Select Seafood' for r in range(a, b + 1))
    check("P-ABS (required choice) names the Modifier Summary rows of that choice + its 'Order' column",
          ok, choice.value if choice else None)
    for label, cell in rows.items():
        check(f'  P-ABS {label}: Details starts with "EDIT →"', str(cell.value).startswith('EDIT →'), cell.value)


def check_x_colour():
    ws = load_workbook('report.xlsx')['Dish Item']
    by_range = {}
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            by_range.setdefault(str(cf.sqref), []).append(rule)
    ok, detail = True, []
    for sqref, rules in by_range.items():
        xr = [r for r in rules if any('="x"' in f for f in r.formula)]
        ar = [r for r in rules if any('<>' in f for f in r.formula)]
        if not ar:
            continue
        colour = ''
        if xr and xr[0].dxf is not None and xr[0].dxf.fill is not None:
            colour = f'{xr[0].dxf.fill.fgColor.rgb}{xr[0].dxf.fill.bgColor.rgb}'
        if not xr or xr[0].priority >= ar[0].priority or '595959' not in colour:
            ok = False
            detail.append(sqref)
    check('Dish Item: "x" rule on every price column, dark grey, ahead of the apricot rule', ok and by_range, detail)


# ── Part 2 ───────────────────────────────────────────────────────────────

def raw_to_tsv(raw, tsv_dir):
    os.makedirs(tsv_dir, exist_ok=True)
    wb = load_workbook(raw)
    for ws in wb.worksheets:
        lines = ['\t'.join('' if v is None else str(v) for v in row)
                 for row in ws.iter_rows(values_only=True)]
        with open(os.path.join(tsv_dir, f'{ws.title}.tsv'), 'w', encoding='utf-8') as f:
            f.write('\n'.join(lines) + '\n')


def tsv_cells(path):
    rows = open(path, encoding='utf-8').read().split('\n')
    return [r.split('\t') for r in rows if r != '']


def edit(path):
    """What a specialist does on the sheet: prices typed both ways, x, strikethrough, notes."""
    wb = load_workbook(path, rich_text=True)
    di = wb['Dish Item']
    hdr = [c.value for c in di[1]]
    rows = {di.cell(r, 2).value: r for r in range(2, di.max_row + 1) if di.cell(r, 2).value}
    col = {h: i + 1 for i, h in enumerate(hdr) if h}
    di.cell(rows['Egg Roll'], col['Order']).value = 'X'          # remove this price
    di.cell(rows['Egg Roll'], col['Small']).value = 'x'          # empty cell — nothing to remove
    di.cell(rows['Buffalo Wings'], col['Order']).value = '$9.25'  # the old '$ habit
    di.cell(rows['Chicken Fingers'], col['Specialty']).value = 8.95
    di.cell(rows['Fried Scallops'], 2).font = Font(strike=True)  # make unavailable
    di.cell(rows['Chicken Wings'], 2).value = CellRichText(
        [TextBlock(InlineFont(strike=True), 'Chicken'), ' Wings'])  # partly struck
    ms = wb['Modifier Summary']
    mh = [c.value for c in ms[1]]
    for r in range(2, ms.max_row + 1):
        if ms.cell(r, 2).value == 'Extra Sauce':
            ms.cell(r, mh.index('Order') + 1).value = 0                # free now
    q = wb['Questionnaire']
    q.cell(2, 9).value = 'Buffalo Wings is now called Buffalo Chicken Wings — rename it'
    q.cell(3, 9).value = 'Egg Roll is gone for good'
    wb.save(path)


def check_part2(scripts):
    raw_to_tsv('raw.xlsx', 'tsv')
    run(scripts, 'wsheet_xlsx_bridge.py', 'to-xlsx', 'tsv', '-o', 'raw_b.xlsx')
    before = {t: tsv_cells(os.path.join('tsv', f'{t}.tsv')) for t in ('dish_item', 'modifier_item', 'modifier_rule')}
    shutil.copy('report.xlsx', 'edited.xlsx')
    edit('edited.xlsx')
    rc_r, out_r = run(scripts, 'specialist_requests.py', 'edited.xlsx', '--dir', '.')
    run(scripts, 'action4_apply_modifier_changes.py', 'edited.xlsx', '-o', 'final_review.xlsx')
    rc5, out5 = run(scripts, 'action5_to_raw_gsheet.py', 'final_review.xlsx', 'raw_b.xlsx', '-o', 'final_raw.xlsx')
    rcp, outp = run(scripts, 'modifier_rule_pricing.py', 'preflight', 'final_raw.xlsx')
    rca, outa = run(scripts, 'modifier_rule_pricing.py', 'apply', 'edited.xlsx', 'final_raw.xlsx', '-o', 'final_raw.xlsx')
    rct, outt = run(scripts, 'wsheet_xlsx_bridge.py', 'to-tsv', 'final_raw.xlsx', 'tsv',
                    '--tabs', 'dish_item,modifier_item,modifier_rule')
    rcrep, outrep = run(scripts, 'part2_report.py', '.', '--restaurant-name', 'Delta Wok')
    check('Part 2 steps run (requests, action5, preflight, apply, to-tsv, report)',
          all(x == 0 for x in (rc_r, rc5, rcp, rca, rct, rcrep)), (rc_r, rc5, rcp, rca, rct, rcrep, outa[-300:]))

    wb = load_workbook('final_raw.xlsx')
    di = wb['dish_item']
    h3 = [c.value for c in di[3]]
    live = {di.cell(r, h3.index('dish_item') + 1).value: r for r in range(4, di.max_row + 1)}

    def val(d, v):
        return di.cell(live[d], h3.index(v) + 1).value
    check('#1 "$9.25" typed the old way → raw $9.25; 8.95 typed plain → raw $8.95',
          val('Buffalo Wings', 'Order') == '$9.25' and val('Chicken Fingers', 'Specialty') == '$8.95',
          (val('Buffalo Wings', 'Order'), val('Chicken Fingers', 'Specialty')))
    mi = wb['modifier_item']
    m3 = [c.value for c in mi[3]]
    es = next((mi.cell(r, m3.index('Order') + 1).value for r in range(4, mi.max_row + 1)
               if mi.cell(r, m3.index('modifier_item') + 1).value == 'Extra Sauce'), None)
    check('#1 0 is a real price: Modifier Summary 0.50 → 0 reaches the raw as $0.00', es == '$0.00', es)
    check('#3 x over Egg Roll Order clears the raw cell', val('Egg Roll', 'Order') is None, val('Egg Roll', 'Order'))
    s = json.load(open('part2_summary.json')) if os.path.isfile('part2_summary.json') else {}
    rem = s.get('action5', {}).get('removed', {'cells': [], 'x_on_empty_cell_ignored': []})
    check('#3 the removal is recorded (was $1.75); x on the empty Egg Roll Small is ignored and reported',
          [(c['dish'], c['variant'], c['old']) for c in rem['cells']] == [('Egg Roll', 'Order', '$1.75')]
          and rem['x_on_empty_cell_ignored'] == [{'dish': 'Egg Roll', 'variant': 'Small'}], rem)

    after = {t: tsv_cells(os.path.join('tsv', f'{t}.tsv')) for t in before}
    changed = []
    for t in before:
        for ri, (o, n) in enumerate(zip(before[t], after[t])):
            for ci in range(max(len(o), len(n))):
                ov = o[ci] if ci < len(o) else ''
                nv = n[ci] if ci < len(n) else ''
                if ov != nv:
                    changed.append((t, ri + 1, before[t][2][ci] if ci < len(before[t][2]) else ci, ov, nv))
    blanked = [c for c in changed if c[4] == '']
    check('#3 the TSV (what wsheet gets) has exactly that one cell cleared', len(blanked) == 1
          and blanked[0][0] == 'dish_item' and blanked[0][3] == '$1.75', blanked)
    check('#1 every price the TSV received is in $X.XX format',
          changed and all(c[4] == '' or re.match(r'^\$\d+\.\d\d$', c[4]) for c in changed),
          [c for c in changed if c[4] and not re.match(r'^\$\d+\.\d\d$', c[4])][:5])

    req = s.get('specialist_requests', {})
    un = req.get('make_unavailable', {}).get('dishes', [])
    check('#2 strikethrough → make permanently unavailable: Fried Scallops (whole) + Chicken Wings (part)',
          sorted((u['dish'], u['partial']) for u in un) == [('Chicken Wings', True), ('Fried Scallops', False)],
          un)
    notes = req.get('notes_for_wonka', {}).get('notes', [])
    check('#7 both notes handed to Wonka, with their Questionnaire row',
          [(n['row'], n['note'][:12]) for n in notes] == [(2, 'Buffalo Wing'), (3, 'Egg Roll is ')], notes)
    check('#2/#7 part2_report shows the removal and what was handed to Wonka',
          'removed' in outrep and 'Egg Roll [Order] was $1.75' in outrep
          and 'make permanently unavailable: 2' in outrep and 'notes (Questionnaire Notes → Wonka): 2' in outrep,
          outrep[-900:])


def older_layout(src, dst):
    """A report as Part 1 wrote it before 2026-09-30: Questionnaire column H was
    "Notes" (the script's own text), with no "Notes → Wonka" column."""
    wb = load_workbook(src)
    q = wb['Questionnaire']
    hdr = [str(c.value or '') for c in q[1]]
    q.cell(1, hdr.index('Details') + 1).value = 'Notes'
    q.delete_cols(next(i + 1 for i, h in enumerate(hdr) if h.startswith('Notes')))
    wb.save(dst)


def check_old_report(scripts, real_report, work, label='older report'):
    d = os.path.join(work, 'old_' + re.sub(r'\W+', '_', label))
    os.makedirs(d, exist_ok=True)
    os.chdir(d)
    shutil.copy(real_report, 'edited.xlsx')
    rc, out = run(scripts, 'specialist_requests.py', 'edited.xlsx', '--dir', '.')
    s = (json.load(open('part2_summary.json')).get('specialist_requests')
         if os.path.isfile('part2_summary.json') else None) or {'notes_for_wonka': {'total': -1}}
    check(f'{label}: its Notes column (script text) is not read as instructions',
          rc == 0 and s['notes_for_wonka']['total'] == 0 and s['notes_for_wonka'].get('skipped_reason'),
          out[-300:])
    rc4, out4 = run(scripts, 'action4_apply_modifier_changes.py', 'edited.xlsx', '-o', 'fr.xlsx')
    check(f'{label}: action4 still reads it', rc4 == 0, out4[-300:])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('bundle')
    ap.add_argument('--real-report', help='an edited report made before 2026-09-30 (optional)')
    a = ap.parse_args()
    real_report = os.path.abspath(a.real_report) if a.real_report else None
    scripts = os.path.join(os.path.abspath(a.bundle), 'menu-change-payload', 'scripts')
    work = tempfile.mkdtemp(prefix='report_checks_')

    print('\nR99999 Tray Test Kitchen (no multi-column pricing, has Party Trays)')
    fixture(work, 'R99999', 'party_tray', [], 'R99999_raw.xlsx', 'R99999_menu_ocr.json')
    part1(scripts, has_sides=False)
    check_numeric('R99999')
    check_questionnaire('R99999', expect_tray=True)

    print('\nR99902 Absolute Wok (multi-column, $0 dishes priced on modifiers)')
    fixture(work, 'R99902', 'modifier_rule', [], 'R99902_raw.xlsx', 'R99902_menu_ocr.json')
    part1(scripts, has_sides=True)
    check_numeric('R99902')
    check_pabs_pointers()

    print('\nR99901 Delta Wok (multi-column) — then a specialist-edited copy through Part 2')
    fixture(work, 'R99901', 'modifier_rule', [], 'R99901_raw.xlsx', 'R99901_menu_ocr.json')
    part1(scripts, has_sides=True)
    check_numeric('R99901')
    check_questionnaire('R99901', expect_tray=False)
    check_x_colour()
    check_part2(scripts)
    print('\nOlder report layout (made from R99901: column H "Notes", no "Notes → Wonka")')
    older_layout(os.path.join(work, 'R99901', 'report.xlsx'), os.path.join(work, 'older_layout.xlsx'))
    check_old_report(scripts, os.path.join(work, 'older_layout.xlsx'), work, 'older layout')

    if real_report:
        print(f'\nOlder edited report: {os.path.basename(real_report)}')
        check_old_report(scripts, real_report, work)

    os.chdir(HERE)
    shutil.rmtree(work, ignore_errors=True)
    print(f'\nreport: {sum(RESULTS)}/{len(RESULTS)} PASS')
    sys.exit(0 if all(RESULTS) else 1)


if __name__ == '__main__':
    main()
