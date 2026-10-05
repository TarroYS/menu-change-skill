"""Pass/fail checks for the maintainer's same-price inference rules (2026-09-30).

    python3 check_rules.py <bundle_dir>

Fixture R99905 "Rules Wok" (no multi-column pricing):
  * General Tso / Orange / Sesame <protein> always cost the same — including a
    Sesame Chicken whose OLD price differed, and a Kung Pao Chicken whose T1
    cohort would say otherwise (the family rule wins); printed siblings that
    disagree (Beef) → no inference.
  * Healthy "Steam X" takes the regular "X" price only where the OLD prices
    were the same (per size), and only when "Steam X" is NOT on the paper menu
    (maintainer, 2026-09-30) — a Healthy dish printed in one size is left alone.
  * Chow Mein and Chop Suey are one family for the same-old-price cohort.
  * a $0.00 freebie is never priced, even when its whole group moved alike.
Exit 1 if any check fails."""
import json
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', 'modifier_rule'))
from build_fixture import make_raw  # noqa: E402

RESULTS = []
PURPLE = 'E4DFEC'


def check(desc, ok, detail=''):
    RESULTS.append(bool(ok))
    print(f"  {'PASS' if ok else 'FAIL'}  {desc}" + ('' if ok or not detail else f'  -- {detail}'))


DISHES = [  # (dish, group, {variant: old price})
    ("General Tso's Chicken", 'Chicken', {'Small': 8.95, 'Large': 12.95}),
    ('Orange Chicken', 'Chicken', {'Small': 8.95, 'Large': 12.95}),
    ('Sesame Chicken', 'Chicken', {'Small': 8.50, 'Large': 12.50}),      # old differs
    ('Kung Pao Chicken', 'Chicken', {'Large': 12.95}),                  # T1 donor, other price
    ('Chicken with Broccoli', 'Chicken', {'Small': 8.00, 'Large': 13.75}),
    ("General Tso's Beef", 'Beef', {'Large': 14.95}),
    ('Orange Beef', 'Beef', {'Large': 14.95}),
    ('Sesame Beef', 'Beef', {'Large': 14.95}),
    ('Beef with Broccoli', 'Beef', {'Small': 9.25, 'Large': 14.75}),
    ('General Tsos Shrimp', 'Seafood', {'Large': 15.95}),                # spelling variant
    ('Sesame Shrimp', 'Seafood', {'Large': 15.95}),
    ('Steam Chicken with Broccoli', 'Healthy Dishes', {'Small': 8.00, 'Large': 13.75}),
    ('Steam Beef with Broccoli', 'Healthy Dishes', {'Small': 9.00, 'Large': 14.75}),
    ('Steam Mixed Vegetables', 'Healthy Dishes', {'Small': 7.00}),     # no regular twin …
    ('Mixed Vegetables with Tofu', 'Vegetables', {'Small': 7.00}),     # … only a near-miss name
    ('Shrimp with Broccoli', 'Seafood', {'Small': 9.50, 'Large': 15.50}),
    ('Steam Shrimp with Broccoli', 'Healthy Dishes', {'Small': 9.50, 'Large': 15.50}),  # printed (Large)
    ('Chicken Chow Mein', 'Chow Mein', {'Small': 7.50, 'Large': 12.25}),
    ('Shrimp Chow Mein', 'Chow Mein', {'Small': 8.00, 'Large': 12.75}),
    ('Chicken Chop Suey', 'Chop Suey', {'Small': 7.50, 'Large': 12.25}),
    ('Shrimp Chop Suey', 'Chop Suey', {'Small': 8.00, 'Large': 12.75}),
    ('House Special Chop Suey', 'Chop Suey', {'Small': 8.50, 'Large': 13.95}),
    ('Egg Roll', 'Appetizers', {'Order': 1.75}),
    ('Spring Roll', 'Appetizers', {'Order': 1.75}),
    ('Crab Rangoon', 'Appetizers', {'Order': 6.25}),
    ('Free Egg Roll', 'Appetizers', {'Order': 0.00}),
]
PAPER = [  # (section, dish, {variant: paper price})
    ('Chicken', "General Tso's Chicken", {'Small': 9.45, 'Large': 13.95}),
    ('Chicken', 'Kung Pao Chicken', {'Large': 13.50}),
    ('Chicken', 'Chicken with Broccoli', {'Small': 8.25, 'Large': 14.25}),
    ('Beef', "General Tso's Beef", {'Large': 15.95}),
    ('Beef', 'Orange Beef', {'Large': 16.25}),                          # siblings disagree
    ('Beef', 'Beef with Broccoli', {'Small': 9.75, 'Large': 15.25}),
    ('Seafood', 'General Tsos Shrimp', {'Large': 16.95}),
    ('Chow Mein', 'Chicken Chow Mein', {'Small': 7.75, 'Large': 12.50}),
    ('Seafood', 'Shrimp with Broccoli', {'Small': 9.95, 'Large': 15.95}),
    ('Healthy Dishes', 'Steam Shrimp with Broccoli', {'Large': 15.95}),     # on paper, Small not
    ('Appetizers', 'Egg Roll', {'Order': 1.95}),
    ('Appetizers', 'Spring Roll', {'Order': 1.95}),
    ('Appetizers', 'Crab Rangoon', {'Order': 6.45}),
    ('Vegetables', 'Mixed Vegetables with Tofu', {'Small': 7.50}),
]


def main():
    bundle = os.path.abspath(sys.argv[1])
    scripts = os.path.join(bundle, 'menu-change-payload', 'scripts')
    work = tempfile.mkdtemp(prefix='pi_rules_')
    os.chdir(work)
    make_raw('raw.xlsx', 'R99905', 'Rules Wok', ['Order', 'Small', 'Large'], DISHES, {}, {}, [])
    sections = {}
    for sec, dish, prices in PAPER:
        sections.setdefault(sec, []).append({'name': dish, 'prices': prices})
    with open('ocr.json', 'w') as f:
        json.dump({'restaurant': {'name': 'Rules Wok'}, 'sections': [
            {'menu_section': s, 'price_columns': [], 'default_price': None, 'items': items}
            for s, items in sections.items()]}, f)

    def run(*a):
        p = subprocess.run([sys.executable, *a], capture_output=True, text=True)
        return p.returncode, p.stdout + p.stderr
    run(os.path.join(scripts, 'action1_raw_to_2a.py'), 'raw.xlsx', '-o', 'preview.xlsx')
    run(os.path.join(scripts, 'action2_ocr_to_2b_0625v2.py'), 'ocr.json', 'preview.xlsx', '-o', 'ocr_out.xlsx')
    run(os.path.join(scripts, 'action3_compare_and_merge.py'), 'preview.xlsx', 'ocr_out.xlsx', '-o', 'report.xlsx')
    rc, out = run(os.path.join(scripts, 'price_inference.py'), 'preview.xlsx', 'ocr_out.xlsx',
                  '--report', 'report.xlsx', '--mode', 'apply')
    check('price_inference runs', rc == 0, out[-400:])

    wb = load_workbook('report.xlsx')
    di = wb['Dish Item']
    hdr = [c.value for c in di[1]]
    cells = {}
    for r in range(2, di.max_row + 1):
        name = di.cell(r, 2).value
        for v in ('Order', 'Small', 'Large'):
            c = di.cell(r, hdr.index(v) + 1)
            if name and c.value not in (None, ''):
                f = c.fill.fgColor.rgb[-6:] if c.fill and c.fill.fill_type else ''
                cells[(name, v)] = (round(float(str(c.value).replace('$', '')), 2), f == PURPLE)
    tiers, reasons = {}, {}
    for row in wb['Price Inference'].iter_rows(min_row=6, values_only=True):
        if row[0] and str(row[0]).startswith('INF-'):
            tiers[(row[4], row[5])] = (row[2], row[9])
            reasons[(row[4], row[5])] = str(row[-1] or '')

    print('\n#4 General Tso / Orange / Sesame always the same price')
    for d in ('Orange Chicken', 'Sesame Chicken'):
        check(f'{d} takes General Tso\'s Chicken\'s printed prices (9.45 / 13.95), purple',
              cells.get((d, 'Small')) == (9.45, True) and cells.get((d, 'Large')) == (13.95, True),
              (cells.get((d, 'Small')), cells.get((d, 'Large')), tiers.get((d, 'Large'))))
    check('the family rule wins over a disagreeing T1 cohort (Kung Pao 13.50 at the same old price)',
          str(tiers.get(('Orange Chicken', 'Large'), ('', ''))[1]).startswith('T0 same-price family'),
          tiers.get(('Orange Chicken', 'Large')))
    check('Sesame Shrimp takes "General Tsos Shrimp" 16.95 (spelling variant matched)',
          cells.get(('Sesame Shrimp', 'Large')) == (16.95, True), cells.get(('Sesame Shrimp', 'Large')))
    check('printed siblings disagree (Gen Tso Beef 15.95 vs Orange Beef 16.25) → Sesame Beef not inferred',
          cells.get(('Sesame Beef', 'Large')) == (14.95, False)
          and tiers.get(('Sesame Beef', 'Large'), ('', ''))[1] == 'T0 conflict',
          (cells.get(('Sesame Beef', 'Large')), tiers.get(('Sesame Beef', 'Large'))))

    print('\n#5 Healthy dish = its regular dish when the old prices were the same')
    check('Steam Chicken with Broccoli takes Chicken with Broccoli 8.25 / 14.25 (old prices equal)',
          cells.get(('Steam Chicken with Broccoli', 'Small')) == (8.25, True)
          and cells.get(('Steam Chicken with Broccoli', 'Large')) == (14.25, True),
          (cells.get(('Steam Chicken with Broccoli', 'Small')),
           cells.get(('Steam Chicken with Broccoli', 'Large'))))
    check('Steam Beef with Broccoli: Large inferred 15.25 (old equal); Small NOT (old 9.00 vs 9.25)',
          cells.get(('Steam Beef with Broccoli', 'Large')) == (15.25, True)
          and cells.get(('Steam Beef with Broccoli', 'Small')) == (9.00, False),
          (cells.get(('Steam Beef with Broccoli', 'Small')), cells.get(('Steam Beef with Broccoli', 'Large'))))
    check('the skip reason names the regular dish and its different old price',
          'regular Beef with Broccoli was $9.25 in Small, not $9.00'
          in reasons.get(('Steam Beef with Broccoli', 'Small'), ''),
          reasons.get(('Steam Beef with Broccoli', 'Small')))
    check('a Healthy dish ON the paper menu is left alone: Steam Shrimp with Broccoli Small stays 9.50 '
          '(its regular twin had the same old price) — the reason says why',
          cells.get(('Steam Shrimp with Broccoli', 'Small')) == (9.50, False)
          and 'on the paper menu' in reasons.get(('Steam Shrimp with Broccoli', 'Small'), ''),
          (cells.get(('Steam Shrimp with Broccoli', 'Small')),
           reasons.get(('Steam Shrimp with Broccoli', 'Small'))))
    check('no regular twin → not inferred (Steam Mixed Vegetables; "Mixed Vegetables with Tofu", same old '
          '$7.00, printed $7.50, is a different dish)',
          cells.get(('Steam Mixed Vegetables', 'Small')) == (7.00, False), cells.get(('Steam Mixed Vegetables', 'Small')))

    print('\n#6 Chow Mein and Chop Suey: same old price → same new price')
    check('Chicken Chop Suey takes Chicken Chow Mein 7.75 / 12.50 (same old prices)',
          cells.get(('Chicken Chop Suey', 'Small')) == (7.75, True)
          and cells.get(('Chicken Chop Suey', 'Large')) == (12.50, True)
          and 'Chow Mein + Chop Suey' in str(tiers.get(('Chicken Chop Suey', 'Large'), ('', ''))[1]),
          (cells.get(('Chicken Chop Suey', 'Small')), tiers.get(('Chicken Chop Suey', 'Large'))))
    check('no printed Chow Mein at the same old price → Shrimp / House Special Chop Suey not inferred',
          cells.get(('Shrimp Chop Suey', 'Small')) == (8.00, False)
          and cells.get(('House Special Chop Suey', 'Large')) == (13.95, False))

    print('\nOther')
    check('the $0.00 freebie is never priced (whole group moved +0.20)',
          cells.get(('Free Egg Roll', 'Order')) == (0.0, False), cells.get(('Free Egg Roll', 'Order')))
    q = [str(row[0]) for row in wb['Questionnaire'].iter_rows(values_only=True) if row[0]]
    check('no P6 section in the Questionnaire (the Price Inference tab + purple cells are the list)',
          not any(x.startswith('P6') for x in q), [x for x in q if x.startswith('P6')])

    shutil.rmtree(work, ignore_errors=True)
    print(f'\nprice_inference rules: {sum(RESULTS)}/{len(RESULTS)} PASS')
    sys.exit(0 if all(RESULTS) else 1)


if __name__ == '__main__':
    main()
