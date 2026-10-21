"""Dish codes (maintainer decisions, 2026-10-05): Part 1's Dish Code tab shows each
size's code next to the code printed on the paper; Part 2 writes a code ONLY where
the specialist put one in a code cell.

    python3 check_codes.py <bundle_dir>

Pins: one printed number on a two-size row → both sizes (DC2); a Lunch / Dinner
code → its own size (DC5); fills and corrections both need the specialist (DC1);
the tab's conventions match prices — x removes, a code in an empty cell adds one
(DC3); a code left on a moved / removed price is only listed (DC4, in check_edits).
Exit 1 if any check fails."""
import json
import os
import shutil
import sys
import tempfile

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', 'modifier_rule'))
sys.path.insert(0, os.path.join(HERE, '..', 'edits'))
sys.path.insert(0, os.path.join(HERE, '..', 'lib'))
from build_fixture import make_raw  # noqa: E402
import check_edits as ce  # noqa: E402
from xlformula import cf_kind  # noqa: E402

RESULTS = []
TAB = 'Dish Code'


def check(desc, ok, detail=''):
    RESULTS.append(bool(ok))
    print(f"  {'PASS' if ok else 'FAIL'}  {desc}" + ('' if ok or not detail else f'  -- {detail}'))


def build():
    variants = ['Order', 'Small', 'Large', 'Lunch', 'Dinner']
    dishes = [('Egg Roll', 'Appetizers', {'Order': 1.75}),
              ('Wonton Soup', 'Soup', {'Small': 3.25, 'Large': 5.25}),
              ('Hot and Sour Soup', 'Soup', {'Small': 3.50, 'Large': 5.50}),
              ('Kung Pao Chicken', 'Chicken', {'Small': 9.25, 'Large': 13.25}),
              ('Sesame Chicken', 'Chicken', {'Small': 9.25, 'Large': 13.25}),
              ('General Tso Chicken', 'Chicken', {'Lunch': 8.95, 'Dinner': 11.95}),
              ('Beef with Broccoli', 'Beef', {'Small': 9.50, 'Large': 13.95}),
              ('Fried Rice', 'Rice', {'Small': 7.00})]
    codes = {'Egg Roll': {'Order': 'A1'}, 'Wonton Soup': {'Small': 'S1', 'Large': 'S2'},
             'Sesame Chicken': {'Small': '15', 'Large': '15'}}
    make_raw('raw.xlsx', 'R99930', 'Code Wok', variants, dishes, {}, {}, [], codes=codes,
             also_in=[('Chef Specials', 'Kung Pao Chicken')])
    sec = lambda name, items, cols=(): {'menu_section': name, 'price_columns': list(cols), 'default_price': None,
                                        'items': items}
    json.dump({'restaurant': {'name': 'Code Wok'}, 'sections': [
        sec('Appetizers', [{'name': 'Egg Roll', 'code': 'A1', 'prices': {'Order': 1.95}}]),
        sec('Soup', [{'name': 'Wonton Soup', 'code': '12', 'prices': {'Small': 3.50, 'Large': 5.50}},
                     {'name': 'Hot and Sour Soup', 'prices': {'Small': 3.75, 'Large': 5.75}}]),
        sec('Chicken', [{'name': 'Kung Pao Chicken', 'code': 'C3', 'prices': {'Small': 9.45, 'Large': 13.45}},
                        {'name': 'Sesame Chicken', 'code': '15', 'prices': {'Small': 9.50, 'Large': 13.50}}]),
        sec('Chef Specials', [{'name': 'Sesame Chicken', 'code': 'S7', 'prices': {'Large': 13.50}}]),
        sec('Dinner Specials', [{'name': 'General Tso Chicken', 'code': 'D26', 'prices': {'Dinner': 12.25}}], ['Dinner']),
        sec('Beef', [{'name': 'Beef with Broccoli', 'prices': {'Small': 9.75, 'Large': 14.25}}]),
        sec('Rice', [{'name': 'Fried Rice', 'prices': {'Small': 7.25}}])],
        'lunch_specials': {'items': [{'name': 'General Tso Chicken', 'code': 'L26', 'price': 9.25}]}},
        open('ocr.json', 'w'))


def tab_cells(path):
    """{(dish, group, column header): (value, fill)} on the Dish Code tab."""
    ws = load_workbook(path)[TAB]
    hdr = [c.value for c in ws[1]]
    out = {}
    for r in range(2, ws.max_row + 1):
        d, g = ws.cell(r, 2).value, ws.cell(r, 1).value
        if not d:
            continue
        for c, h in enumerate(hdr, 1):
            if h and not str(h).startswith('_ref_') and c > 3:
                cell = ws.cell(r, c)
                out[(d, g, h)] = (cell.value, cell.fill.fgColor.rgb[-6:] if cell.fill.fill_type else '')
    return out


def live_codes(path):
    ws = load_workbook(path)['dish_item']
    h2 = [str(c.value or '') for c in ws[2]]
    h3 = [c.value for c in ws[3]]
    out = {}
    for r in range(4, ws.max_row + 1):
        if str(ws.cell(r, 1).value or '').startswith('{'):
            out[ws.cell(r, 2).value] = {h3[i]: ws.cell(r, i + 1).value for i, m in enumerate(h2)
                                        if m.endswith('unit_code') and ws.cell(r, i + 1).value not in (None, '')}
    return out


def set_live_code(path, dish, size, value):
    wb = load_workbook(path)
    ws = wb['dish_item']
    h2 = [str(c.value or '') for c in ws[2]]
    h3 = [c.value for c in ws[3]]
    ci = next(i + 1 for i, m in enumerate(h2) if m.endswith('unit_code') and h3[i] == size)
    r = next(r for r in range(4, ws.max_row + 1)
             if ws.cell(r, 2).value == dish and str(ws.cell(r, 1).value or '').startswith('{'))
    ws.cell(r, ci).value = value
    wb.save(path)


def edit(src, dst, values, sort=False):
    """Type values into the Dish Code tab: {(dish, group, size): value}. sort: then sort
    the visible block by dish name, descending, like a quick sort in a spreadsheet
    (the hidden baselines stay where they are)."""
    wb = load_workbook(src)
    ws = wb[TAB]
    hdr = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(hdr) if h}
    rows = {(ws.cell(r, 2).value, ws.cell(r, 1).value): r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value}
    for (d, g, size), v in values.items():
        ws.cell(rows[(d, g)], col[size]).value = v
    if sort:
        vis = next(i for i, h in enumerate(hdr) if str(h or '').startswith('_ref_'))
        data = [[ws.cell(r, c).value for c in range(1, vis + 1)] for r in range(2, ws.max_row + 1)]
        data.sort(key=lambda row: str(row[1] or ''), reverse=True)
        for i, row in enumerate(data, 2):
            for c, v in enumerate(row, 1):
                ws.cell(i, c).value = v
    wb.save(dst)
    return wb


def ocr(sections, lunch=None, name='Wok'):
    js = {'restaurant': {'name': name}, 'sections': [
        {'menu_section': s, 'price_columns': list(cols), 'default_price': None, 'items': items}
        for s, cols, items in sections]}
    if lunch:
        js['lunch_specials'] = {'items': lunch}
    json.dump(js, open('ocr.json', 'w'))


def p_code(path='report.xlsx'):
    """The Questionnaire's P-CODE rows, as text."""
    return [' '.join(str(c) for c in row if c not in (None, ''))
            for row in load_workbook(path)['Questionnaire'].iter_rows(values_only=True)
            if str(row[0] or '').startswith('P-CODE-')]


def set_cached_text(path, sheet, coord, value):
    """Give a formula cell a saved text result, as Excel / Google Sheets would."""
    import re
    import zipfile
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    wbx = items['xl/workbook.xml'].decode()
    tag = re.search(r'<sheet\b[^>]*\bname="%s"[^>]*/>' % re.escape(sheet), wbx).group(0)
    rid = re.search(r'r:id="([^"]+)"', tag).group(1)
    rels = items['xl/_rels/workbook.xml.rels'].decode()
    rel = re.search(r'<Relationship\b[^>]*\bId="%s"[^>]*/>' % re.escape(rid), rels).group(0)
    target = re.search(r'Target="([^"]+)"', rel).group(1).lstrip('/')
    target = target if target.startswith('xl/') else 'xl/' + target
    x = items[target].decode()
    x2 = re.sub(r'<c r="%s"([^>]*?)>\s*(<f>[^<]*</f>)\s*(?:<v\s*/>|<v>[^<]*</v>)?\s*</c>' % coord,
                lambda m: '<c r="%s"%s t="str">%s<v>%s</v></c>' % (
                    coord, re.sub(r'\s+t="[^"]*"', '', m.group(1)), m.group(2), value), x)
    assert x2 != x, coord
    items[target] = x2.encode()
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, d in items.items():
            z.writestr(n, d)


def review_part1(scripts, work):
    """What the paper proposes (second review, 2026-10-05)."""
    print('\nPart 1 — which printed codes are proposed, and where')
    d = os.path.join(work, 'r1')
    os.makedirs(d)
    os.chdir(d)
    make_raw('raw.xlsx', 'R99933', 'Review Wok', ['Order', 'Small', 'Large', 'Lunch', 'Dinner'],
             [('Fried Dumplings', 'Appetizers', {'Order': 8.50}),
              ('Crab Rangoon', 'Appetizers', {'Order': 6.25}),
              ('Spring Roll', 'Appetizers', {'Order': 1.75}),
              ('Chicken Lo Mein', 'Lo Mein', {'Small': 7.25, 'Large': 10.25}),
              ('Pork Lo Mein', 'Lo Mein', {'Small': 7.25, 'Large': 10.25}),
              ('Mongolian Beef', 'Beef', {'Small': 9.25, 'Large': 13.25}),
              ('Sesame Chicken', 'Chicken', {'Order': 12.95}),
              ('General Tso Chicken', 'Chicken', {'Lunch': 8.95, 'Dinner': 11.95}),
              ('Orange Chicken', 'Chicken', {'Lunch': 8.95, 'Dinner': 11.95}),
              ('Kung Pao Chicken', 'Chicken', {'Lunch': 8.95, 'Dinner': 11.95}),
              ('Hunan Chicken', 'Chicken', {'Lunch': 8.95, 'Dinner': 11.95})], {}, {}, [],
             codes={'Fried Dumplings': {'Order': 'A7'}, 'Spring Roll': {'Order': '春1'},
                    'Sesame Chicken': {'Order': '26'}})
    ocr([('Appetizers', [], [{'name': '10 Fried Dumplings', 'prices': {'Order': 8.95}},      # a count
                             {'name': 'A5 Crab Rangoon', 'prices': {'Order': 6.50}},
                             {'name': 'Spring Roll', 'code': '夏1', 'prices': {'Order': 1.95}}]),
         ('Lo Mein', [], [{'name': '12. Chicken or Pork Lo Mein', 'prices': {'Small': 7.50, 'Large': 10.50}}]),
         ('Beef', [], [{'name': 'Mongolian Beef', 'code': '35', 'prices': {'Order': 13.45}}]),   # → Large
         ("Chef's Specials", [], [{'name': 'Mongolian Beef', 'code': 'S8', 'prices': {'Large': 13.50}}]),
         ('Chicken', [], [{'name': 'General Tso Chicken', 'code': '26', 'prices': {'Small': 9.25}}])],   # → Lunch
        lunch=[{'name': 'General Tso Chicken', 'code': 'L26', 'price': 9.50},
               {'name': 'Sesame Chicken', 'code': 'L27', 'price': 8.95}])
    out = ce.part1(scripts)
    check('Part 1 runs', all(v[0] == 0 for v in out.values()), {k: v[1][-300:] for k, v in out.items() if v[0]})
    t = tab_cells('report.xlsx')
    P = ' · paper'
    paper = {(k[0], k[2][:-len(P)]): v for k, v in t.items() if k[2].endswith(P) and v[0]}
    check('"10 Fried Dumplings" proposes no code (10 is a count); "A5 Crab Rangoon" and "12. Chicken or Pork '
          'Lo Mein" do (a letter / a "." after the number)',
          ('Fried Dumplings', 'Order') not in paper and paper.get(('Crab Rangoon', 'Order'), ('',))[0] == 'A5'
          and all(paper.get((dish, s), ('',))[0] == '12' for dish in ('Chicken Lo Mein', 'Pork Lo Mein')
                  for s in ('Small', 'Large')), paper)
    check('a lunch-special code never lands on another size (Sesame Chicken has only Order: L27 not proposed)',
          ('Sesame Chicken', 'Order') not in paper, paper.get(('Sesame Chicken', 'Order')))
    pc = p_code()
    check('two codes reaching one size through different paper sizes are not proposed, but asked: Mongolian '
          'Beef Large (35 via Order, S8) and General Tso Chicken Lunch (26 via Small, the lunch special L26)',
          ('Mongolian Beef', 'Large') not in paper and ('General Tso Chicken', 'Lunch') not in paper
          and any('Mongolian Beef [Large]' in r and '35 / S8' in r for r in pc)
          and any('General Tso Chicken [Lunch]' in r and '26 / L26' in r for r in pc), (pc, paper))
    check('a printed code in another script is compared as itself (system 春1, paper 夏1: blue, differs)',
          paper.get(('Spring Roll', 'Order')) == ('夏1', 'ADD8E6'), paper.get(('Spring Roll', 'Order')))

    d2 = os.path.join(work, 'r1b')
    os.makedirs(d2)
    os.chdir(d2)
    make_raw('raw.xlsx', 'R99934', 'Conflict Wok', ['Order', 'Small', 'Large'],
             [('Mongolian Beef', 'Beef', {'Small': 9.25, 'Large': 13.25})], {}, {}, [])
    ocr([('Beef', [], [{'name': 'Mongolian Beef', 'code': '35', 'prices': {'Order': 13.45}}]),
         ("Chef's Specials", [], [{'name': 'Mongolian Beef', 'code': 'S8', 'prices': {'Large': 13.50}}])])
    ce.part1(scripts)
    check('a restaurant with no codes but two printed for one size still gets the Dish Code tab and the P-CODE '
          'question', TAB in load_workbook('report.xlsx').sheetnames
          and any('Mongolian Beef [Large]' in r for r in p_code()), p_code())


def review_part2(scripts, work):
    """What Part 2 writes from odd cells (second review, 2026-10-05)."""
    print('\nPart 2 — pasted columns, formulas, spaces, numbers, grey cells, follow-ups')
    base = os.path.join(work, 'r2')
    os.makedirs(base)
    os.chdir(base)
    make_raw('raw.xlsx', 'R99935', 'Paste Wok', ['Order', 'Small', 'Large'],
             [('Egg Roll', 'Appetizers', {'Order': 1.75}),
              ('Spring Roll', 'Appetizers', {'Order': 1.75}),
              ('Fried Rice', 'Rice', {'Order': 7.00}),
              ('Combo Plate', 'Combos', {'Order': 9.95}),
              ('Crab Rangoon', 'Appetizers', {'Order': 6.25}),
              ('Shrimp Toast', 'Appetizers', {'Order': 5.25}),
              ('Shrimp Roll', 'Appetizers', {'Order': 2.25}),
              ('Fried Wonton', 'Appetizers', {'Order': 4.25}),
              ('Hot Soup', 'Soup', {'Small': 3.50, 'Large': 5.50}),
              ('Egg Drop Soup', 'Soup', {'Small': 3.25, 'Large': 5.25})], {}, {}, [],
             codes={'Egg Roll': {'Order': 'A1'}, 'Combo Plate': {'Order': 'Combo 9'}, 'Crab Rangoon': {'Order': '1.0'},
                    'Shrimp Toast': {'Order': 'X'}, 'Shrimp Roll': {'Order': 'S  5'}, 'Fried Wonton': {'Order': 'W1'},
                    'Hot Soup': {'Small': 'H1S', 'Large': 'H1L'}, 'Egg Drop Soup': {'Small': 'E1S', 'Large': 'E1L'}})
    ocr([('Appetizers', [], [{'name': 'Egg Roll', 'code': 'A9', 'prices': {'Order': 1.95}},
                             {'name': 'Spring Roll', 'code': 'A2', 'prices': {'Order': 1.95}}])])
    ce.part1(scripts)

    # (a) the whole 'Order · paper' column pasted over 'Order' — header and all
    d = os.path.join(work, 'r2_paste')
    shutil.copytree(base, d)
    os.chdir(d)
    wb = load_workbook('report.xlsx')
    ws = wb[TAB]
    hdr = [c.value for c in ws[1]]
    src_c, dst_c = hdr.index('Order · paper') + 1, hdr.index('Order') + 1
    for r in range(1, ws.max_row + 1):
        ws.cell(r, dst_c).value = ws.cell(r, src_c).value
    wb.save('edited.xlsx')
    p2 = ce.part2(scripts)
    c = live_codes('final_raw.xlsx') if os.path.exists('final_raw.xlsx') else {}
    check('a whole paper column pasted over its code column, header included, is read (Egg Roll A1 → A9, '
          'Spring Roll + A2)', c.get('Egg Roll') == {'Order': 'A9'} and c.get('Spring Roll') == {'Order': 'A2'},
          (c.get('Egg Roll'), c.get('Spring Roll'), p2['a5'][1][-200:]))
    # (b) a code column whose header was changed some other way stops Part 2
    d = os.path.join(work, 'r2_renamed')
    shutil.copytree(base, d)
    os.chdir(d)
    wb = load_workbook('report.xlsx')
    ws = wb[TAB]
    ws.cell(1, [c.value for c in ws[1]].index('Order') + 1).value = 'Order code'
    wb.save('edited.xlsx')
    p2 = ce.part2(scripts)
    check('a code column that cannot be found (header renamed) stops Part 2 with a message — no draft',
          p2['a5'][0] == 1 and 'no code column' in p2['a5'][1] and not os.path.exists('final_raw.xlsx'),
          p2['a5'][1][-300:])

    # (c) odd cells
    d = os.path.join(work, 'r2_cells')
    shutil.copytree(base, d)
    os.chdir(d)
    wb = load_workbook('report.xlsx')
    ws = wb[TAB]
    hdr = [x.value for x in ws[1]]
    colm = {h: i + 1 for i, h in enumerate(hdr) if h}
    rowm = {ws.cell(r, 2).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value}
    at = lambda dish, size: ws.cell(rowm[dish], colm[size])
    at('Egg Roll', 'Order').value = '=' + at('Egg Roll', 'Order · paper').coordinate   # a formula → A9
    lit = at('Spring Roll', 'Order')
    lit.value = '=G3'
    lit.data_type = 's'                          # typed into a text cell: the text "=G3"
    at('Fried Rice', 'Order').value = 'R1\xa0'   # a no-break space pasted along
    at('Fried Wonton', 'Order').value = 'W1\xa0'  # the system code + a no-break space: untouched
    at('Combo Plate', 'Order').value = 'Combo  9'    # the system code with two spaces: untouched
    at('Crab Rangoon', 'Order').value = '1'      # system "1.0" → "1"
    at('Egg Roll', 'Small').value = 'B9'         # grey: Egg Roll has no Small price
    at('Spring Roll', 'Small').value = 'x'       # x where there is no code
    at('Shrimp Roll', 'Order').value = 'x'       # live "S  5" (two spaces) removed
    at('Hot Soup', 'Small').value = 'x'          # its Small price is removed too (Dish Item) — code handled
    di = wb['Dish Item']
    col, _r, rows = ce.report_rows(di)
    for dish in ('Hot Soup', 'Egg Drop Soup'):
        di.cell(rows[(dish, 'Soup')], col['Small']).value = 'x'      # Egg Drop Soup's code is left on Small
    wb.save('edited.xlsx')
    set_cached_text('edited.xlsx', TAB, at('Egg Roll', 'Order').coordinate, 'A9')    # last: openpyxl drops it
    wbv = load_workbook('edited.xlsx')
    kind = lambda dish, size: cf_kind(wbv, wbv[TAB], at(dish, size).coordinate, ce.rule_kind)
    seen = {'=G3': kind('Spring Roll', 'Order'), 'nbsp': kind('Fried Rice', 'Order'),
            'two spaces': kind('Combo Plate', 'Order'), 'W1 + nbsp': kind('Fried Wonton', 'Order'),
            '1.0→1': kind('Crab Rangoon', 'Order'),
            'X untouched': kind('Shrimp Toast', 'Order'), 'grey': kind('Egg Roll', 'Small')}
    check('the tab shows each the way Part 2 treats it: "=G3" red, a no-break space new, two spaces / a trailing '
          'no-break space untouched, '
          '1.0 → 1 an edit, a system code "X" untouched, a code on a grey cell red',
          seen == {'=G3': 'INVALID', 'nbsp': 'NEW', 'two spaces': None, 'W1 + nbsp': None, '1.0→1': 'EDIT',
                   'X untouched': None,
                   'grey': 'INVALID'}, seen)
    p2 = ce.part2(scripts)
    c = live_codes('final_raw.xlsx') if os.path.exists('final_raw.xlsx') else {}
    want = {'Egg Roll': {'Order': 'A9'}, 'Spring Roll': {}, 'Fried Rice': {'Order': 'R1'},
            'Combo Plate': {'Order': 'Combo 9'}, 'Fried Wonton': {'Order': 'W1'}, 'Crab Rangoon': {'Order': '1'},
            'Shrimp Toast': {'Order': 'X'},
            'Shrimp Roll': {}, 'Hot Soup': {'Large': 'H1L'}, 'Egg Drop Soup': {'Small': 'E1S', 'Large': 'E1L'}}
    got = {k: c.get(k) for k in want}
    check('Part 2 writes the formula\'s value (A9), "R1" without the no-break space, "1"; refuses "=G3" and the '
          'grey-cell code; leaves "Combo 9" and "X"', got == want, {k: (got[k], want[k]) for k in want if got[k] != want[k]})
    fr = load_workbook('final_raw.xlsx')['dish_item'] if os.path.exists('final_raw.xlsx') else None
    check('… as text, never a formula', fr is not None
          and all(cell.data_type != 'f' for row in fr.iter_rows(min_row=4) for cell in row),
          fr and [cell.coordinate for row in fr.iter_rows(min_row=4) for cell in row if cell.data_type == 'f'][:5])
    a5 = p2['summary'].get('action5', {})
    ign = {(x['name'], x['variant']): x['why'] for x in a5.get('ignored_input', {}).get('cells', []) if x.get('tab') == TAB}
    fol = [(f['dish'], f['code']) for f in a5.get('code_followups', {}).get('items', [])]
    check('refusals reported ("=G3" not a code, B9 needs a price); a code the specialist removed with its price '
          '(Hot Soup H1S) is not a follow-up, one left on a removed price (Egg Drop Soup E1S) is',
          'leading =' in ign.get(('Spring Roll', 'Order'), '') and 'needs one' in ign.get(('Egg Roll', 'Small'), '')
          and fol == [('Egg Drop Soup', 'E1S')], (ign, fol))
    rep = p2['rep'][1]
    check('part2_report: an x where there is no code says so', 'Spring Roll [Small] — that cell has no code to remove' in rep,
          [ln for ln in rep.splitlines() if 'x on' in ln])
    check('to-tsv applies the fill and the x on "S  5" (codes compared as Part 2 reads them) — exit 0',
          p2['tsv'][0] == 0, p2['tsv'][1][-400:])


def main():
    bundle = os.path.abspath(sys.argv[1])
    scripts = os.path.join(bundle, 'menu-change-payload', 'scripts')
    work = tempfile.mkdtemp(prefix='code_checks_')

    print('\nPart 1 — the Dish Code tab')
    d = os.path.join(work, 'p1')
    os.makedirs(d)
    os.chdir(d)
    build()
    out = ce.part1(scripts)
    check('Part 1 runs', all(v[0] == 0 for v in out.values()), {k: v[1][-300:] for k, v in out.items() if v[0]})
    wb = load_workbook('report.xlsx')
    check('the report has a Dish Code tab, after the other tabs', TAB in wb.sheetnames
          and wb.sheetnames.index(TAB) > wb.sheetnames.index('Dish Item'), wb.sheetnames)
    t = tab_cells('report.xlsx')
    P = ' · paper'
    check('a printed code equal to the system one is green (Egg Roll A1)',
          t[('Egg Roll', 'Appetizers', 'Order')][0] == 'A1' and t[('Egg Roll', 'Appetizers', 'Order' + P)] == ('A1', 'C6EFCE'))
    check('DC2: one number printed for a two-size dish goes on both sizes (Wonton Soup "12" on Small and Large, '
          'blue: the system has S1 / S2)',
          t[('Wonton Soup', 'Soup', 'Small' + P)] == ('12', 'ADD8E6') and t[('Wonton Soup', 'Soup', 'Large' + P)] == ('12', 'ADD8E6')
          and t[('Wonton Soup', 'Soup', 'Small')][0] == 'S1', {k: v for k, v in t.items() if k[0] == 'Wonton Soup'})
    check('a code the system lacks is yellow, on both rows of a dish in two groups (Kung Pao Chicken C3)',
          all(t[('Kung Pao Chicken', g, s + P)] == ('C3', 'FFFF00') for g in ('Chicken', 'Chef Specials')
              for s in ('Small', 'Large')))
    check('DC5: a Lunch-section code goes on Lunch, a Dinner one on Dinner (General Tso Chicken L26 / D26)',
          t[('General Tso Chicken', 'Chicken', 'Lunch' + P)][0] == 'L26'
          and t[('General Tso Chicken', 'Chicken', 'Dinner' + P)][0] == 'D26')
    check('two different codes printed for one size (Sesame Chicken Large: 15 and S7) are not proposed …',
          t[('Sesame Chicken', 'Chicken', 'Large' + P)][0] is None and t[('Sesame Chicken', 'Chicken', 'Small' + P)][0] == '15')
    q = [[c for c in row if c not in (None, '')] for row in wb['Questionnaire'].iter_rows(values_only=True)]
    pcode = [r for r in q if r and str(r[0]).startswith('P-CODE-')]
    check('… they are asked (P-CODE), and a P-CODE row points at the tab for the printed codes to review',
          any('Sesame Chicken [Large]' in map(str, r) and '15 / S7' in ' '.join(map(str, r)) for r in pcode)
          and any('see the Dish Code tab' in ' '.join(map(str, r)) for r in pcode), pcode)
    check('a size without a price is grey (a code needs a price)',
          t[('Egg Roll', 'Appetizers', 'Small')][1] == 'EDEDED' and t[('Egg Roll', 'Appetizers', 'Order')][1] == '')
    check('codes stay text through price_inference / compare (Sesame Chicken "15" is not the number 15)',
          t[('Sesame Chicken', 'Chicken', 'Small')][0] == '15', type(t[('Sesame Chicken', 'Chicken', 'Small')][0]))

    print('\nPart 2 — only what the specialist confirmed is written')
    typed = {('Wonton Soup', 'Soup', 'Small'): '12',                 # accept a printed code (a correction)
             ('Wonton Soup', 'Soup', 'Large'): 12,                   # … typed as a number
             ('General Tso Chicken', 'Chicken', 'Lunch'): 'L26',     # a fill
             ('Egg Roll', 'Appetizers', 'Order'): 'x',               # remove A1
             ('Kung Pao Chicken', 'Chicken', 'Small'): 'C3',         # two rows, two different codes …
             ('Kung Pao Chicken', 'Chef Specials', 'Small'): 'C4',
             ('Kung Pao Chicken', 'Chicken', 'Large'): 'C3',         # … one row typed, the other untouched
             ('Sesame Chicken', 'Chicken', 'Large'): '#',            # not a code
             ('Beef with Broccoli', 'Beef', 'Lunch'): 'B9',          # no price on Lunch
             ('Hot and Sour Soup', 'Soup', 'Small'): 'H1'}           # wsheet sets H0 there since Part 1
    for sort in (False, True):
        tag = 'sorted' if sort else 'plain'
        dd = os.path.join(work, 'p2_' + tag)
        shutil.copytree(d, dd)
        os.chdir(dd)
        wb = edit('report.xlsx', 'edited.xlsx', typed, sort=sort)
        set_live_code('raw.xlsx', 'Hot and Sour Soup', 'Small', 'H0')
        p2 = ce.part2(scripts)
        c = live_codes('final_raw.xlsx')
        want = {'Wonton Soup': {'Small': '12', 'Large': '12'}, 'General Tso Chicken': {'Lunch': 'L26'},
                'Egg Roll': {}, 'Kung Pao Chicken': {'Large': 'C3'}, 'Sesame Chicken': {'Small': '15', 'Large': '15'},
                'Beef with Broccoli': {}, 'Hot and Sour Soup': {'Small': 'H0'}}
        got = {k: c.get(k) for k in want}
        check(f'[{tag}] corrections, fills and x are written; a conflict, "#", a code on a size with no price '
              f'and a code changed in wsheet since are not', got == want, {k: (got[k], want[k]) for k in want if got[k] != want[k]})
        if not sort:
            check('untouched printed codes are never written (General Tso Chicken Dinner D26; Kung Pao Chicken '
                  'Large on the Chef Specials row)', 'Dinner' not in c.get('General Tso Chicken', {}))
            a5 = p2['summary'].get('action5', {})
            ign = {(x['name'], x['variant']): x['why'] for x in a5.get('ignored_input', {}).get('cells', [])
                   if x.get('tab') == TAB}
            conf = [x for x in a5.get('duplicates', {}).get('conflicts', []) if x.get('tab') == TAB]
            check('each refusal is reported: "#" not a code, B9 needs a price, Kung Pao Chicken Small conflict, '
                  'Hot and Sour Soup drift', 'not a dish code' in ign.get(('Sesame Chicken', 'Large'), '')
                  and 'needs one' in ign.get(('Beef with Broccoli', 'Lunch'), '')
                  and any(x['name'] == 'Kung Pao Chicken' for x in conf)
                  and any(x['dish'] == 'Hot and Sour Soup' for x in a5.get('codes', {}).get('drift', {}).get('items', [])),
                  (ign, conf))
            diff = {(r, cc): (o, n) for _t, _x, r, cc, o, n in p2['tsv_diff']}
            check('the TSV gets exactly those code cells (2 corrections, 2 fills, 1 removal)',
                  sorted(v for v in diff.values() if v[1] in ('12', 'L26', 'C3', '') and not str(v[0]).startswith('$'))
                  == sorted([('S1', '12'), ('S2', '12'), ('', 'L26'), ('', 'C3'), ('A1', '')]), diff)
            rep = p2['rep'][1]
            check('part2_report counts them: 2 filled, 2 corrected, 1 removed',
                  '2 filled, 2 corrected, 1 removed' in rep and 'Egg Roll [Order] code A1 → removed' in rep, rep[-800:])
            # the sheet shows each typed cell the way Part 2 treats it
            ws = wb[TAB]
            hdr = [x.value for x in ws[1]]
            colm = {h: i + 1 for i, h in enumerate(hdr) if h}
            rowm = {(ws.cell(r, 2).value, ws.cell(r, 1).value): r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value}
            kind = lambda dish, g, size: cf_kind(wb, ws, ws.cell(rowm[(dish, g)], colm[size]).coordinate, ce.rule_kind)
            seen = {'x': kind('Egg Roll', 'Appetizers', 'Order'), '#': kind('Sesame Chicken', 'Chicken', 'Large'),
                    'new': kind('General Tso Chicken', 'Chicken', 'Lunch'), 'edit': kind('Wonton Soup', 'Soup', 'Small'),
                    'number': kind('Wonton Soup', 'Soup', 'Large'), 'untouched': kind('Sesame Chicken', 'Chicken', 'Small')}
            check('the tab colours each the way Part 2 treats it: x grey, "#" red, a new code yellow, a change '
                  'apricot (a typed number too), untouched none',
                  seen == {'x': 'X', '#': 'INVALID', 'new': 'NEW', 'edit': 'EDIT', 'number': 'EDIT', 'untouched': None}, seen)

    print('\nBridge — an x on a code clears the TSV only while it still holds that code')
    os.chdir(os.path.join(work, 'p2_plain'))
    ce.raw_to_tsv('raw.xlsx', 'tsv2')
    rows_ = ce.tsv_cells('tsv2/dish_item.tsv')
    h2, h3 = rows_[1], rows_[2]
    ci = next(i for i, m in enumerate(h2) if m.endswith('unit_code') and h3[i] == 'Order')
    ri = next(i for i, r in enumerate(rows_) if len(r) > 1 and r[1] == 'Egg Roll' and r[0].startswith('{'))
    rows_[ri][ci] = 'A9'                                    # wsheet changed the code since
    open('tsv2/dish_item.tsv', 'w', encoding='utf-8').write('\n'.join('\t'.join(r) for r in rows_) + '\n')
    rc, out = ce.run(scripts, 'wsheet_xlsx_bridge.py', 'to-tsv', 'final_raw.xlsx', 'tsv2')
    after = ce.tsv_cells('tsv2/dish_item.tsv')
    check('a code changed in the TSV since to-xlsx is not cleared, and to-tsv exits 1',
          after[ri][ci] == 'A9' and rc == 1 and 'NOT applied' in out, out[-300:])

    print('\nA restaurant with no codes anywhere')
    d0 = os.path.join(work, 'nocodes')
    os.makedirs(d0)
    os.chdir(d0)
    make_raw('raw.xlsx', 'R99931', 'Plain Wok', ['Order', 'Small', 'Large'],
             [('Egg Roll', 'Appetizers', {'Order': 1.75})], {}, {}, [])
    json.dump({'restaurant': {'name': 'Plain Wok'}, 'sections': [
        {'menu_section': 'Appetizers', 'price_columns': [], 'default_price': None,
         'items': [{'name': 'Egg Roll', 'prices': {'Order': 1.95}}]}]}, open('ocr.json', 'w'))
    ce.part1(scripts)
    wb = load_workbook('report.xlsx')
    shutil.copy('report.xlsx', 'edited.xlsx')
    p2 = ce.part2(scripts)
    check('no Dish Code tab, no P-CODE row, and part2_report says no codes were written',
          TAB not in wb.sheetnames
          and not any(str(r[0] or '').startswith('P-CODE') for r in wb['Questionnaire'].iter_rows(values_only=True))
          and 'none (no Dish Code tab' in p2['rep'][1], wb.sheetnames)

    print('\nSystem codes, but none printed on the paper')
    d1 = os.path.join(work, 'syscodes')
    os.makedirs(d1)
    os.chdir(d1)
    make_raw('raw.xlsx', 'R99932', 'Sys Code Wok', ['Order', 'Small', 'Large'],
             [('Egg Roll', 'Appetizers', {'Order': 1.75})], {}, {}, [], codes={'Egg Roll': {'Order': 'A1'}})
    json.dump({'restaurant': {'name': 'Sys Code Wok'}, 'sections': [
        {'menu_section': 'Appetizers', 'price_columns': [], 'default_price': None,
         'items': [{'name': 'Egg Roll', 'prices': {'Order': 1.95}}]}]}, open('ocr.json', 'w'))
    ce.part1(scripts)
    ws = load_workbook('report.xlsx')[TAB]
    hdr = [c.value for c in ws[1]]
    check('the tab shows the system codes (editable) and no empty paper block',
          'Order' in hdr and not any(str(h or '').endswith(' · paper') for h in hdr)
          and any(c.value == 'A1' for row in ws.iter_rows(min_row=2) for c in row[3:4]), hdr)

    review_part1(scripts, work)
    review_part2(scripts, work)

    os.chdir(HERE)
    shutil.rmtree(work, ignore_errors=True)
    print(f'\ncodes: {sum(RESULTS)}/{len(RESULTS)} PASS')
    sys.exit(0 if all(RESULTS) else 1)


if __name__ == '__main__':
    main()
